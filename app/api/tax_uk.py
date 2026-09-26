"""UK Making Tax Digital (roadmap 2026-09 §3.6): company MTD settings and the
VAT return boxes 1–9 (app/services/uk_mtd/)."""
from __future__ import annotations

import json
from datetime import date

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.db.session import get_db
from app.services.audit_service import log_audit_event
from app.services.uk_mtd import categories as C
from app.services.uk_mtd import itsa
from app.services.uk_mtd import periods as P
from app.services.uk_mtd import settings as S
from app.services.uk_mtd.vat import vat_return

router = APIRouter(prefix="/tax/uk", tags=["tax"])


def _base_currency(db: Session) -> str:
    from app.services.fx_service import _current_company_row
    row = _current_company_row(db)
    return ((row.base_currency if row is not None else None) or "GBP").upper()


class MtdSettingsPayload(BaseModel):
    income_source: str | None = None
    period_basis: str | None = None
    vat_registered: bool | None = None
    vat_stagger: str | None = None
    vrn: str | None = None
    category_overrides: dict[str, str] | None = None


@router.get("/settings")
def get_settings(db: Session = Depends(get_db)) -> dict:
    return S.get_settings(db)


@router.put("/settings")
def put_settings(payload: MtdSettingsPayload, db: Session = Depends(get_db)) -> dict:
    changes = payload.model_dump(exclude_unset=True)
    try:
        cur = S.save_settings(db, **changes)
    except S.MtdSettingsError as e:
        raise HTTPException(status_code=422, detail=str(e)) from e
    log_audit_event(db, action="update", entity_type="uk_mtd_settings", entity_id="company",
                    detail=json.dumps({k: v for k, v in changes.items() if k != "category_overrides"}
                                      | ({"category_overrides": len(changes["category_overrides"] or {})}
                                         if "category_overrides" in changes else {}), default=str))
    db.commit()
    return cur


@router.get("/vat/periods")
def vat_periods(db: Session = Depends(get_db)) -> dict:
    conf = S.get_settings(db)
    today = date.today()
    rows = [p.as_dict() | {"days_left": (p.deadline - today).days, "open": p.end >= today}
            for p in P.vat_periods_before(today, conf["vat_stagger"], count=6)]
    return {"vat_registered": conf["vat_registered"], "stagger": conf["vat_stagger"], "periods": rows}


def _period(db: Session, period_end: date) -> P.VatPeriod:
    try:
        return P.vat_period_ending(period_end, S.get_settings(db)["vat_stagger"])
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e)) from e


@router.get("/vat/return")
def get_vat_return(period_end: date = Query(...), db: Session = Depends(get_db)) -> dict:
    conf = S.get_settings(db)
    out = vat_return(db, _period(db, period_end), currency=_base_currency(db))
    return out | {"vrn": conf["vrn"], "vat_registered": conf["vat_registered"]}


_BOX_LABELS = {
    "1": "VAT due on sales and other outputs",
    "2": "VAT due on acquisitions from the EU (Northern Ireland)",
    "3": "Total VAT due (box 1 + box 2)",
    "4": "VAT reclaimed on purchases and other inputs",
    "5": "Net VAT to pay to HMRC or reclaim",
    "6": "Total value of sales and other outputs, excluding VAT",
    "7": "Total value of purchases and other inputs, excluding VAT",
    "8": "Total value of dispatches of goods to the EU (Northern Ireland), excluding VAT",
    "9": "Total value of acquisitions of goods from the EU (Northern Ireland), excluding VAT",
}


@router.get("/vat/return/export")
def export_vat_return(period_end: date = Query(...), db: Session = Depends(get_db)) -> Response:
    """The nine boxes as CSV, for bridging software or the records."""
    import csv
    import io
    conf = S.get_settings(db)
    out = vat_return(db, _period(db, period_end), currency=_base_currency(db))
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["VRN", conf["vrn"] or ""])
    w.writerow(["Period", out["period"]["start"], out["period"]["end"]])
    w.writerow(["Due", out["period"]["deadline"]])
    w.writerow([])
    w.writerow(["Box", "Description", f"Amount ({out['currency']})"])
    for n in map(str, range(1, 10)):
        v = out["boxes"][n]
        w.writerow([n, _BOX_LABELS[n], f"{v:.2f}" if int(n) <= 5 else str(int(v))])
    return Response(content=buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="vat-return-{out["period"]["end"]}.csv"'})


# --- MTD for Income Tax (quarterly updates) -------------------------------------------------

def _tax_year(label: str | None) -> int:
    if not label:
        return P.tax_year_of(date.today())
    try:
        return P.parse_tax_year(label)
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e)) from e


@router.get("/itsa/quarters")
def itsa_quarters(tax_year: str | None = None, db: Session = Depends(get_db)) -> dict:
    conf = S.get_settings(db)
    year = _tax_year(tax_year)
    today = date.today()
    quarters = [q.as_dict() | {"days_left": (q.deadline - today).days, "open": q.end >= today}
                for q in P.itsa_quarters(year, conf["period_basis"])]
    out = {"tax_year": P.tax_year_label(year), "basis": conf["period_basis"], "source": conf["income_source"],
           "quarters": quarters}
    if conf["income_source"] in C.CATALOGUE:
        out["mandation"] = itsa.mandation(db, year, source=conf["income_source"],
                                          overrides=conf["category_overrides"], currency=_base_currency(db))
    return out


def _update(db: Session, tax_year: str | None, quarter: int) -> dict:
    conf = S.get_settings(db)
    if conf["income_source"] not in C.CATALOGUE:
        raise HTTPException(status_code=409, detail="Choose the income source (self-employment or UK property) "
                                                    "in the Making Tax Digital settings first.")
    q = P.ItsaQuarter(_tax_year(tax_year), quarter, conf["period_basis"])
    return itsa.quarterly_update(db, q, source=conf["income_source"], overrides=conf["category_overrides"],
                                 currency=_base_currency(db))


@router.get("/itsa/update")
def itsa_update(quarter: int = Query(..., ge=1, le=4), tax_year: str | None = None,
                db: Session = Depends(get_db)) -> dict:
    return _update(db, tax_year, quarter)


@router.get("/itsa/categories")
def itsa_categories(db: Session = Depends(get_db)) -> dict:
    """Every income/expense account with its HMRC category (default or
    override) for the chosen source, and the categories to choose from."""
    from sqlalchemy import select

    from app.models.account import Account
    from app.services.reporting.common import EXPENSE, REVENUE, classify_account_code
    conf = S.get_settings(db)
    source = conf["income_source"]
    accounts = []
    if source in C.CATALOGUE:
        for a in db.execute(select(Account).order_by(Account.code)).scalars():
            if classify_account_code(a.code) not in (REVENUE, EXPENSE) or len(a.code) < 4:
                continue
            accounts.append({"code": a.code, "name": a.name,
                             "default": C.default_category(a.code, source),
                             "override": conf["category_overrides"].get(a.code)})
    return {"source": source, "catalogue": {k: {"income": list(v["income"]), "expenses": list(v["expenses"])}
                                            for k, v in C.CATALOGUE.items()},
            "excluded": C.EXCLUDED, "accounts": accounts}


@router.get("/itsa/update/export")
def itsa_export(quarter: int = Query(..., ge=1, le=4), tax_year: str | None = None,
                db: Session = Depends(get_db)) -> Response:
    """Workbook: the update by HMRC category, the accounts behind it, and the
    request body for HMRC's cumulative update."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font
    up = _update(db, tax_year, quarter)
    wb = Workbook()
    ws = wb.active
    ws.title = "Update"
    bold = Font(bold=True)
    q = up["quarter"]
    for row in (("Tax year", q["tax_year"]), ("Quarter", q["quarter"]), ("Basis", q["basis"]),
                ("Quarter", f"{q['start']} to {q['end']}"), ("Cumulative from", q["cumulative_start"]),
                ("Due", q["deadline"]), ("Income source", up["source"]), ("Currency", up["currency"]), ()):
        ws.append(list(row))
    ws.append(["Category", "Kind", "This quarter", "Tax year to date"])
    for c in ws[ws.max_row]:
        c.font = bold
    for kind in ("income", "expenses"):
        for line in up[kind]:
            ws.append([line["category"], kind, line["quarter"], line["year_to_date"]])
    ws.append([])
    for label, key in (("Profit this quarter", "quarter"), ("Profit tax year to date", "year_to_date")):
        ws.append([label, "", up["totals"][key]["profit"]])
    ws.column_dimensions["A"].width = 34
    acc = wb.create_sheet("Accounts")
    acc.append(["Code", "Account", "HMRC category", "Kind", "This quarter", "Tax year to date", "Overridden"])
    for c in acc[1]:
        c.font = bold
    for a in up["accounts"]:
        acc.append([a["code"], a["name"], a["category"], a["kind"], a["quarter"], a["year_to_date"],
                    "yes" if a["overridden"] else ""])
    for a in up["excluded"]:
        acc.append([a["code"], a["name"], "excluded", "", "", a["amount"], ""])
    acc.column_dimensions["B"].width = 40
    acc.column_dimensions["C"].width = 28
    body = wb.create_sheet("HMRC body")
    body.append(["Cumulative update request body (itemised)"])
    body.append([json.dumps(up["hmrc_body"], indent=2)])
    if up["hmrc_body_consolidated"]:
        body.append(["Consolidated expenses alternative"])
        body.append([json.dumps(up["hmrc_body_consolidated"], indent=2)])
    buf = io.BytesIO()
    wb.save(buf)
    name = f"mtd-itsa-{q['tax_year']}-q{q['quarter']}.xlsx"
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})
