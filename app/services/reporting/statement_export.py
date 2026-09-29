"""The financial statements as a PDF or an Excel workbook, made on the server (roadmap §4.9).

The statement services already return presentation-ordered rows — the
Iranian and UK templates as flat ``rows`` (``row_type``, ``indent_level``,
current / prior / opening amounts), changes in equity as a component matrix,
and the generic statements as sections of account trees. ``tables`` turns
whichever the company's locale uses into one ``Table`` shape; ``render_pdf``
prints them one per page under the company's letterhead (right to left, Jalali
dates and Persian digits for an Iranian company) and ``render_xlsx`` writes
one sheet each, with real numbers, bold totals and indented labels.

Deductions the statements mark ``is_negative_presentation`` come out negative
in the workbook — so a column still adds up — and in parentheses in the PDF.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date
from typing import Any

from sqlalchemy.orm import Session

ALL = ("balance_sheet", "income_statement", "comprehensive_income", "changes_in_equity", "cash_flow")
GENERIC = ("balance_sheet", "income_statement", "cash_flow")

TITLES = {
    "ir": {
        "balance_sheet": ("صورت وضعیت مالی", "Statement of financial position"),
        "income_statement": ("صورت سود و زیان", "Income statement"),
        "comprehensive_income": ("صورت سود و زیان جامع", "Statement of comprehensive income"),
        "changes_in_equity": ("صورت تغییرات در حقوق مالکانه", "Statement of changes in equity"),
        "cash_flow": ("صورت جریان‌های نقدی", "Statement of cash flows"),
    },
    "uk": {
        "balance_sheet": ("", "Balance sheet"),
        "income_statement": ("", "Profit and loss account"),
        "comprehensive_income": ("", "Statement of comprehensive income"),
        "changes_in_equity": ("", "Statement of changes in equity"),
        "cash_flow": ("", "Statement of cash flows"),
    },
    "default": {
        "balance_sheet": ("ترازنامه", "Balance sheet"),
        "income_statement": ("صورت سود و زیان", "Income statement"),
        "cash_flow": ("صورت جریان وجوه نقد", "Cash flow statement"),
    },
}
WORDS = {
    "fa": {"as_of": "در تاریخ", "period": "از {a} تا {b}", "total": "جمع", "item": "شرح", "prior": "دوره قبل",
           "opening": "ابتدای دوره قبل", "page": "صفحه", "amounts": "مبالغ به {ccy}", "prepared": "تهیه شده در {d}"},
    "en": {"as_of": "As at", "period": "{a} to {b}", "total": "Total", "item": "", "prior": "Prior period",
           "opening": "Prior period opening", "page": "Page", "amounts": "Amounts in {ccy}", "prepared": "Prepared {d}"},
}
STYLES = {"line", "subtotal", "total", "header", "spacer"}


@dataclass
class Table:
    key: str
    title: str
    subtitle: str
    columns: list[str]
    rows: list[dict] = field(default_factory=list)   # {label, indent, style, values: [int|None]}


class ExportError(ValueError):
    pass


# --- building the statements --------------------------------------------------------------------------------------

def _locale(db: Session) -> str:
    from app.services.locale_service import get_reporting_locale
    loc = (get_reporting_locale(db) or "default").strip().lower()
    return loc if loc in ("ir", "uk") else "default"


def available(db: Session) -> tuple[str, tuple[str, ...]]:
    loc = _locale(db)
    return loc, (GENERIC if loc == "default" else ALL)


def _fmt_date(d, locale: str, lang: str) -> str:
    from app.services.documents.formatting import fmt_date
    if isinstance(d, str):                                 # the balance sheets carry ISO strings
        try:
            d = date.fromisoformat(d[:10])
        except ValueError:
            return d
    return fmt_date(d, "ir" if lang == "fa" and locale == "ir" else "uk")


def _label(row, lang: str) -> str:
    if hasattr(row, "label_fa"):
        return (row.label_fa if lang == "fa" else (row.label_en or row.label_fa)) or ""
    return getattr(row, "label", "") or ""


def _flat(resp, lang: str, *, three: bool) -> list[dict]:
    out = []
    for r in resp.rows:
        deduction = bool(getattr(r, "is_negative_presentation", False))
        vals = [r.amount_current, r.amount_prior] + ([getattr(r, "amount_prior_beginning", None)] if three else [])
        vals = [None if v is None else (-abs(int(v)) if deduction else int(v)) for v in vals]
        style = r.row_type if r.row_type in STYLES else "line"
        if style in ("header", "spacer"):
            vals = [None] * len(vals)
        out.append({"label": _label(r, lang), "indent": int(r.indent_level or 0), "style": style, "values": vals})
    return out


def _matrix(resp, lang: str) -> tuple[list[str], list[dict]]:
    comps = list(resp.components)
    cols = [(_label(c, lang)) for c in comps] + [WORDS[lang]["total"]]
    rows = []
    for r in resp.rows:
        cells = {c.component: c.amount for c in r.cells}
        style = r.row_type if r.row_type in STYLES else "line"
        vals = [None if cells.get(c.key) is None else int(cells.get(c.key) or 0) for c in comps] + [
            None if r.total is None else int(r.total)]
        rows.append({"label": _label(r, lang), "indent": 0, "style": style,
                     "values": [None] * len(vals) if style in ("header", "spacer") else vals})
    return cols, rows


def _sections(resp, lang: str) -> list[dict]:
    """Generic statements: each section's account tree, then its total."""
    out: list[dict] = []

    def walk(node, depth):
        name = (node.label_fa if lang == "fa" and getattr(node, "label_fa", None) else node.account_name)
        out.append({"label": f"{node.account_code} {name}".strip(), "indent": depth + 1,
                    "style": "line", "values": [int(node.balance or 0)]})
        for ch in node.children or []:
            walk(ch, depth + 1)

    for sec in resp.sections.values():
        out.append({"label": sec.label_fa if lang == "fa" else sec.label, "indent": 0, "style": "header", "values": [None]})
        items = getattr(sec, "items", None)
        if items is not None:
            for node in items:
                walk(node, 0)
            total = sec.total
        else:                                                      # cash-flow sections
            for ln in sec.lines:
                out.append({"label": f"{ln.account_code} {(ln.label_fa if lang == 'fa' and ln.label_fa else ln.account_name)}",
                            "indent": 1, "style": "line", "values": [int(ln.amount or 0)]})
            total = sec.net
        out.append({"label": (WORDS[lang]["total"] + " " + (sec.label_fa if lang == "fa" else sec.label)).strip(),
                    "indent": 0, "style": "subtotal", "values": [int(total or 0)]})
    for k, v in (resp.totals or {}).items():
        out.append({"label": k.replace("_", " ").capitalize(), "indent": 0, "style": "total", "values": [int(v or 0)]})
    return out


def tables(db: Session, *, statements: list[str] | None = None, from_date: date | None = None,
           to_date: date | None = None, currency: str | None = None, lang: str | None = None) -> tuple[list[Table], dict]:
    """The chosen statements (default: every one the locale has), in order."""
    from app.services.reporting.cash_flow_service import CashFlowService
    from app.services.reporting.financial_statement_service import FinancialStatementService
    from app.services.reporting.iran_statement_service import IranStatementService
    from app.services.reporting.uk_statement_service import UKStatementService

    locale, have = available(db)
    lang = lang if lang in ("fa", "en") else ("fa" if locale == "ir" else "en")
    if lang == "fa" and locale == "uk":
        lang = "en"                                                # the UK template has English labels only
    wanted = [s for s in (statements or have)]
    unknown = [s for s in wanted if s not in have]
    if unknown:
        raise ExportError(f"Not a statement for this company: {', '.join(unknown)}. Choose from {', '.join(have)}.")
    today = date.today()
    to_date = to_date or today
    from_date = from_date or date(to_date.year, 1, 1)
    if from_date > to_date:
        raise ExportError("The period starts after it ends.")
    W = WORDS[lang]
    d = lambda x: _fmt_date(x, locale, lang)  # noqa: E731
    period_cols = [W["period"].format(a=d(from_date), b=d(to_date)), W["prior"]]
    out: list[Table] = []
    title = lambda key: TITLES[locale][key][0 if lang == "fa" else 1] or TITLES[locale][key][1]  # noqa: E731
    flows = {"from_date": from_date, "to_date": to_date, "currency": currency}

    for key in wanted:
        if locale == "ir":
            svc = IranStatementService(db)
            if key == "balance_sheet":
                r = svc.balance_sheet(as_of=to_date, currency=currency)
                cols = [d(r.as_of), d(r.comparative_as_of) if r.comparative_as_of else W["prior"],
                        d(r.comparative_beginning_as_of) if r.comparative_beginning_as_of else W["opening"]]
                out.append(Table(key, title(key), f"{W['as_of']} {d(to_date)}", cols, _flat(r, lang, three=True)))
                continue
            r = {"income_statement": svc.income_statement, "comprehensive_income": svc.comprehensive_income,
                 "cash_flow": svc.cash_flow, "changes_in_equity": svc.changes_in_equity}[key](**flows)
        elif locale == "uk":
            svc = UKStatementService(db)
            if key == "balance_sheet":
                r = svc.balance_sheet(as_of=to_date, currency=currency)
                cols = [d(r.as_of), d(r.comparative_as_of) if r.comparative_as_of else W["prior"]]
                out.append(Table(key, title(key), f"{W['as_of']} {d(to_date)}", cols, _flat(r, lang, three=False)))
                continue
            r = {"income_statement": svc.income_statement, "comprehensive_income": svc.comprehensive_income,
                 "cash_flow": svc.cash_flow, "changes_in_equity": svc.changes_in_equity}[key](**flows)
        else:
            if key == "balance_sheet":
                r = FinancialStatementService(db).balance_sheet(to_date=to_date, currency=currency)
                out.append(Table(key, title(key), f"{W['as_of']} {d(to_date)}", [d(to_date)], _sections(r, lang)))
            elif key == "income_statement":
                r = FinancialStatementService(db).income_statement(**flows)
                out.append(Table(key, title(key), period_cols[0], [period_cols[0]], _sections(r, lang)))
            else:
                r = CashFlowService(db).statement(**flows)
                out.append(Table(key, title(key), period_cols[0], [period_cols[0]], _sections(r, lang)))
            continue
        if key == "changes_in_equity":
            cols, rows = _matrix(r, lang)
            out.append(Table(key, title(key), period_cols[0], cols, rows))
        else:
            out.append(Table(key, title(key), period_cols[0], period_cols, _flat(r, lang, three=False)))
    from app.services.fx_service import get_reporting_currency
    meta = {"locale": locale, "lang": lang, "from_date": from_date, "to_date": to_date,
            "currency": (currency or get_reporting_currency(db) or "").upper()}
    return out, meta


# --- PDF -----------------------------------------------------------------------------------------------------------

def _num(v, lang: str) -> str:
    from app.services.documents.formatting import to_persian_digits
    if v is None:
        return ""
    text = f"{abs(int(v)):,}"
    text = f"({text})" if int(v) < 0 else text
    return to_persian_digits(text) if lang == "fa" else text


def render_pdf(db: Session, tabs: list[Table], meta: dict, *, cover: dict | None = None,
               extra_html: str = "") -> bytes:
    from app.services.documents.branding import build_brand
    from app.services.documents.engine import render_pdf as _render
    brand = build_brand(db)
    lang = meta["lang"]
    rtl = lang == "fa"
    W = WORDS[lang]
    ctx = {
        "rtl": rtl, "dir": "rtl" if rtl else "ltr", "lang": lang, "brand": brand, "brand_color": brand["brand_color"],
        "font_family": brand["font_family"], "words": W, "cover": cover, "extra_html": extra_html,
        "amounts_in": W["amounts"].format(ccy=meta.get("currency") or ""),
        "prepared": W["prepared"].format(d=_fmt_date(date.today(), meta["locale"], lang)),
        "tables": [{"title": t.title, "subtitle": t.subtitle, "columns": t.columns,
                    "rows": [{**r, "cells": [_num(v, lang) for v in r["values"]]} for r in t.rows]} for t in tabs],
    }
    return _render(ctx, "statements.html")


# --- Excel ---------------------------------------------------------------------------------------------------------

def render_xlsx(tabs: list[Table], meta: dict, *, company: str = "", extra_sheets: list[tuple[str, list[list[Any]]]] | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    rtl = meta["lang"] == "fa"
    wb = Workbook()
    wb.remove(wb.active)
    money = "#,##0;(#,##0);-"
    thin = Side(style="thin", color="94A3B8")
    for t in tabs:
        # Excel allows 31 characters; "Statement of …" sheets go by their subject
        name = t.title[len("Statement of "):].capitalize() if t.title.startswith("Statement of ") else t.title
        ws = wb.create_sheet(name[:31].replace("/", "-") or t.key[:31])
        ws.sheet_view.rightToLeft = rtl
        ws.append([company])
        ws.append([t.title])
        ws.append([t.subtitle + "  ·  " + WORDS[meta["lang"]]["amounts"].format(ccy=meta.get("currency") or "")])
        ws.append([])
        ws["A1"].font = Font(bold=True, size=12)
        ws["A2"].font = Font(bold=True, size=14)
        ws["A3"].font = Font(italic=True, color="64748B")
        header = [WORDS[meta["lang"]]["item"]] + t.columns
        ws.append(header)
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=5, column=c)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F1F5F9")
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for r in t.rows:
            if r["style"] == "spacer":
                ws.append([])
                continue
            ws.append([r["label"]] + [v for v in r["values"]])
            row = ws.max_row
            label = ws.cell(row=row, column=1)
            label.alignment = Alignment(indent=min(int(r["indent"]) * 2, 15))
            bold = r["style"] in ("subtotal", "total", "header")
            for c in range(1, len(header) + 1):
                cell = ws.cell(row=row, column=c)
                if bold:
                    cell.font = Font(bold=True)
                if c > 1:
                    cell.number_format = money
                if r["style"] == "total":
                    cell.border = Border(top=thin, bottom=Side(style="double", color="0F172A"))
                elif r["style"] == "subtotal":
                    cell.border = Border(top=thin)
        ws.column_dimensions["A"].width = 52
        for c in range(2, len(header) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 20
        ws.freeze_panes = "B6"
    for name, rows in extra_sheets or []:
        ws = wb.create_sheet(name[:31])
        ws.sheet_view.rightToLeft = rtl
        for row in rows:
            ws.append(row)
        if rows:
            for cell in ws[1]:
                cell.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(meta: dict, kind: str, ext: str) -> str:
    return f"{kind}-{meta['from_date']}-{meta['to_date']}.{ext}"


# the month helper the close pack uses
def month_range(key: str) -> tuple[date, date]:
    from app.services.calendar_periods import key_bounds
    try:
        return key_bounds(key)
    except Exception as e:  # noqa: BLE001
        raise ExportError(f"'{key}' isn't a month (YYYY-MM).") from e

