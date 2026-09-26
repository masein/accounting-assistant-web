"""Iranian seasonal tax filings (roadmap 2026-09 §3.2).

* **گزارش معاملات فصلی (ماده ۱۶۹ مکرر, TTMS)** — every counterparty the
  company bought from or sold to in a Jalali season, with the identity the
  tax organisation asks for, the amount before tax and the VAT, due **45 days**
  after the season ends. Sales already registered in سامانه مودیان are left
  out by default: an electronic invoice accepted there needs no second report.
* **اظهارنامه مالیات بر ارزش افزوده** — output VAT on sales, input VAT on
  purchases, returns (reduction credit notes) netted, the balance payable or
  carried forward; due **15 days** after the season ends.

Both are computed from the same recognised invoices (issued, partially paid,
paid; never drafts or cancelled), with each line's VAT worked out exactly as
it was when the invoice posted, so the report, the return and the ledger
agree to the rial. Only rial (IRR) invoices are reported; others are listed
so their rial equivalent can be entered by hand. Deadlines are not moved for
Fridays or public holidays — the tax organisation extends those to the next
working day.
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal
from typing import Any

import jdatetime
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models.credit_note import CreditNote
from app.models.entity import Entity
from app.models.invoice import Invoice
from app.services.tax_service import _line_tax

SEASONS = {1: "بهار", 2: "تابستان", 3: "پاییز", 4: "زمستان"}
TTMS_DAYS = 45
VAT_DAYS = 15
RECOGNISED = ("issued", "partially_paid", "paid")
REPORT_CURRENCY = "IRR"
# What TTMS needs from every counterparty before the file is accepted.
REQUIRED_PARTY_FIELDS = ("national_id", "economic_code", "postal_code", "address", "phone")
FIELD_LABELS = {
    "national_id": "شناسه/کد ملی", "economic_code": "کد اقتصادی", "postal_code": "کد پستی",
    "address": "نشانی", "phone": "تلفن", "entity": "طرف معامله",
}


# --- seasons -------------------------------------------------------------------------

@dataclass(frozen=True)
class Season:
    year: int      # Jalali year, e.g. 1405
    season: int    # 1 بهار … 4 زمستان

    @property
    def name(self) -> str:
        return f"{SEASONS[self.season]} {self.year}"

    @property
    def start(self) -> date:
        return jdatetime.date(self.year, 3 * self.season - 2, 1).togregorian()

    @property
    def end(self) -> date:
        month = 3 * self.season
        last = 31 if month <= 6 else 30
        if month == 12:
            last = 30 if jdatetime.date(self.year, 1, 1).isleap() else 29
        return jdatetime.date(self.year, month, last).togregorian()

    @property
    def ttms_deadline(self) -> date:
        return self.end + timedelta(days=TTMS_DAYS)

    @property
    def vat_deadline(self) -> date:
        return self.end + timedelta(days=VAT_DAYS)

    def previous(self) -> "Season":
        return Season(self.year - 1, 4) if self.season == 1 else Season(self.year, self.season - 1)

    def as_dict(self) -> dict[str, Any]:
        def j(d: date) -> str:
            return jdatetime.date.fromgregorian(date=d).strftime("%Y/%m/%d")
        return {
            "year": self.year, "season": self.season, "name": self.name,
            "start": self.start.isoformat(), "end": self.end.isoformat(),
            "start_jalali": j(self.start), "end_jalali": j(self.end),
            "ttms_deadline": self.ttms_deadline.isoformat(), "ttms_deadline_jalali": j(self.ttms_deadline),
            "vat_deadline": self.vat_deadline.isoformat(), "vat_deadline_jalali": j(self.vat_deadline),
        }


def season_of(d: date) -> Season:
    jd = jdatetime.date.fromgregorian(date=d)
    return Season(jd.year, (jd.month - 1) // 3 + 1)


def parse_season(year: int, season: int) -> Season:
    if not (1300 <= int(year) <= 1600) or int(season) not in SEASONS:
        raise ValueError("year must be a Jalali year (e.g. 1405) and season 1–4")
    return Season(int(year), int(season))


# --- amounts -----------------------------------------------------------------------------

def _share(part: int, whole: int, amount: int) -> int:
    """amount × part / whole, half-up to the rial."""
    if not whole:
        return 0
    return int((Decimal(amount) * Decimal(part) / Decimal(whole)).to_integral_value(rounding=ROUND_HALF_UP))


def invoice_split(inv: Invoice) -> dict[str, int]:
    """Base split by VAT treatment and the VAT, line by line exactly as the
    invoice posted (app/api/invoices._tax_breakdown)."""
    items = list(inv.items or [])
    if not items:  # an invoice with no lines carries no VAT
        amount = int(inv.amount or 0)
        return {"taxable_base": 0, "exempt_base": amount, "base": amount, "vat": 0}
    taxable = exempt = vat = 0
    for it in items:
        total = int(it.line_total or 0)
        tax = _line_tax(total, it.tax_rate, it.taxable)
        treatment = getattr(it, "tax_treatment", "standard") or "standard"
        if it.taxable and float(it.tax_rate or 0) > 0 and treatment == "standard":
            taxable += total
            vat += tax
        else:
            exempt += total
    return {"taxable_base": taxable, "exempt_base": exempt, "base": taxable + exempt, "vat": vat}


_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def ascii_digits(value: str | None) -> str:
    """Identity numbers go to TTMS in Latin digits, however they were typed."""
    return (value or "").strip().translate(_DIGITS)


def person_type(ent: Entity | None) -> str:
    """حقیقی (10-digit national code) / حقوقی (11-digit national id), as the
    مودیان builder reads it; unknown until one of them is entered."""
    if ent is None:
        return "نامشخص"
    nid = "".join(ch for ch in ascii_digits(ent.national_id) if ch.isdigit())
    econ = "".join(ch for ch in ascii_digits(ent.economic_code) if ch.isdigit())
    if len(nid) == 10 or len(econ) == 14:
        return "حقیقی"
    if len(nid) == 11 or len(econ) == 11:
        return "حقوقی"
    return "نامشخص"


def missing_fields(ent: Entity | None) -> list[str]:
    if ent is None:
        return ["entity"]
    return [f for f in REQUIRED_PARTY_FIELDS if not (getattr(ent, f, None) or "").strip()]


# --- the season's invoices -----------------------------------------------------------------

def _season_invoices(db: Session, s: Season) -> list[Invoice]:
    return db.execute(
        select(Invoice).where(Invoice.issue_date >= s.start, Invoice.issue_date <= s.end,
                              Invoice.status.in_(RECOGNISED), Invoice.kind.in_(("sales", "purchase")))
        .options(selectinload(Invoice.items)).order_by(Invoice.issue_date, Invoice.number)
    ).scalars().all()


def _season_returns(db: Session, s: Season) -> list[CreditNote]:
    return db.execute(
        select(CreditNote).where(CreditNote.date >= s.start, CreditNote.date <= s.end,
                                 CreditNote.note_type == "reduction")
        .order_by(CreditNote.date)
    ).scalars().all()


def _return_kind(cn: CreditNote, inv: Invoice | None) -> str:
    """A reduction note belongs to the side of its invoice; a standalone one
    to the side it was recorded on."""
    if inv is not None and inv.kind in ("sales", "purchase"):
        return inv.kind
    return "purchase" if (cn.kind or "sales") == "purchase" else "sales"


def _jalali(d: date) -> str:
    return jdatetime.date.fromgregorian(date=d).strftime("%Y/%m/%d")


def quarterly_report(db: Session, s: Season, *, include_moadian: bool = False) -> dict[str, Any]:
    """TTMS rows per counterparty for sales and purchases, the invoices behind
    them, and every counterparty whose identity is incomplete."""
    invoices = _season_invoices(db, s)
    returns = _season_returns(db, s)
    by_id = {inv.id: inv for inv in invoices}
    ent_ids = {inv.entity_id for inv in invoices if inv.entity_id} | {cn.entity_id for cn in returns if cn.entity_id}
    entities = {e.id: e for e in db.execute(select(Entity).where(Entity.id.in_(ent_ids))).scalars()} if ent_ids else {}

    other_currency: list[dict[str, Any]] = []
    excluded_moadian: dict[str, int] = {"count": 0, "base": 0, "vat": 0}
    groups: dict[tuple[str, Any], dict[str, Any]] = {}
    details: list[dict[str, Any]] = []

    def group(kind: str, entity_id) -> dict[str, Any]:
        key = (kind, entity_id)
        if key not in groups:
            ent = entities.get(entity_id)
            groups[key] = {
                "kind": kind, "entity_id": str(entity_id) if entity_id else None,
                "person_type": person_type(ent),
                "national_id": ascii_digits(ent.national_id) if ent else "",
                "economic_code": ascii_digits(ent.economic_code) if ent else "",
                "name": ((ent.legal_name or ent.name) if ent else "بدون طرف معامله"),
                "postal_code": ascii_digits(ent.postal_code) if ent else "",
                "province": (ent.province or "") if ent else "", "city": (ent.city or "") if ent else "",
                "address": (ent.address or "") if ent else "", "phone": ascii_digits(ent.phone) if ent else "",
                "invoice_count": 0, "base": 0, "vat": 0, "total": 0,
                "returns_base": 0, "returns_vat": 0,
                "missing": missing_fields(ent),
            }
        return groups[key]

    for inv in invoices:
        split = invoice_split(inv)
        in_moadian = inv.kind == "sales" and inv.moadian_status == "confirmed"
        detail = {
            "kind": inv.kind, "number": inv.number, "issue_date": inv.issue_date.isoformat(),
            "issue_date_jalali": _jalali(inv.issue_date),
            "counterparty": (entities[inv.entity_id].legal_name or entities[inv.entity_id].name)
            if inv.entity_id in entities else None,
            "currency": inv.currency or REPORT_CURRENCY, "base": split["base"], "vat": split["vat"],
            "total": split["base"] + split["vat"], "moadian_status": inv.moadian_status,
            "reported": True,
        }
        if (inv.currency or REPORT_CURRENCY).upper() != REPORT_CURRENCY:
            detail["reported"] = False
            other_currency.append({"number": inv.number, "currency": inv.currency, "total": detail["total"],
                                   "kind": inv.kind})
        elif in_moadian and not include_moadian:
            detail["reported"] = False
            excluded_moadian["count"] += 1
            excluded_moadian["base"] += split["base"]
            excluded_moadian["vat"] += split["vat"]
        else:
            g = group(inv.kind, inv.entity_id)
            g["invoice_count"] += 1
            g["base"] += split["base"]
            g["vat"] += split["vat"]
        details.append(detail)

    for cn in returns:
        inv = by_id.get(cn.invoice_id) if cn.invoice_id else None
        if inv is None and cn.invoice_id:
            inv = db.get(Invoice, cn.invoice_id)
        if (cn.currency or REPORT_CURRENCY).upper() != REPORT_CURRENCY:
            continue
        if inv is not None and inv.kind == "sales" and inv.moadian_status == "confirmed" and not include_moadian:
            continue  # the return of a مودیان invoice goes to مودیان as a return invoice
        # The VAT inside a return is the credited share of its invoice's VAT.
        if inv is not None:
            split = invoice_split(inv)
            gross = split["base"] + split["vat"]
            vat = _share(split["vat"], gross, int(cn.amount or 0))
        else:
            vat = 0
        g = group(_return_kind(cn, inv), cn.entity_id or (inv.entity_id if inv else None))
        g["returns_vat"] += vat
        g["returns_base"] += int(cn.amount or 0) - vat

    rows = {"sales": [], "purchase": []}
    for g in groups.values():
        g["total"] = g["base"] + g["vat"]
        g["net_base"] = g["base"] - g["returns_base"]
        g["net_vat"] = g["vat"] - g["returns_vat"]
        rows[g["kind"]].append(g)
    for kind in rows:
        rows[kind].sort(key=lambda r: (-r["total"], r["name"]))

    def totals(kind: str) -> dict[str, int]:
        rs = rows[kind]
        return {k: sum(r[k] for r in rs) for k in ("invoice_count", "base", "vat", "total", "returns_base",
                                                   "returns_vat", "net_base", "net_vat")}

    incomplete = [{"kind": r["kind"], "name": r["name"], "entity_id": r["entity_id"],
                   "missing": r["missing"], "missing_labels": [FIELD_LABELS[m] for m in r["missing"]]}
                  for kind in rows for r in rows[kind] if r["missing"]]
    return {
        "season": s.as_dict(),
        "include_moadian": include_moadian,
        "sales": rows["sales"], "purchases": rows["purchase"],
        "totals": {"sales": totals("sales"), "purchases": totals("purchase")},
        "excluded_moadian": excluded_moadian,
        "other_currency": other_currency,
        "incomplete": incomplete,
        "details": details,
        "ready": not incomplete and not other_currency,
        "note": "Sales accepted in سامانه مودیان are left out unless included. Purchases backed by an "
                "electronic invoice you confirmed in the taxpayer portal can also be removed before filing.",
    }


def vat_return(db: Session, s: Season) -> dict[str, Any]:
    """Figures for the seasonal VAT return, from every recognised rial invoice
    (مودیان or not) and the season's reduction credit notes."""
    invoices = _season_invoices(db, s)
    out = {side: {"taxable_base": 0, "exempt_base": 0, "vat": 0, "returns_base": 0, "returns_vat": 0,
                  "invoice_count": 0} for side in ("sales", "purchases")}
    left_out = 0
    for inv in invoices:
        if (inv.currency or REPORT_CURRENCY).upper() != REPORT_CURRENCY:
            left_out += 1
            continue
        side = out["sales" if inv.kind == "sales" else "purchases"]
        split = invoice_split(inv)
        side["taxable_base"] += split["taxable_base"]
        side["exempt_base"] += split["exempt_base"]
        side["vat"] += split["vat"]
        side["invoice_count"] += 1
    for cn in _season_returns(db, s):
        if (cn.currency or REPORT_CURRENCY).upper() != REPORT_CURRENCY:
            continue
        inv = db.get(Invoice, cn.invoice_id) if cn.invoice_id else None
        vat = 0
        if inv is not None:
            split = invoice_split(inv)
            vat = _share(split["vat"], split["base"] + split["vat"], int(cn.amount or 0))
        side = out["sales" if _return_kind(cn, inv) == "sales" else "purchases"]
        side["returns_vat"] += vat
        side["returns_base"] += int(cn.amount or 0) - vat
    for side in out.values():
        side["net_vat"] = side["vat"] - side["returns_vat"]
    payable = out["sales"]["net_vat"] - out["purchases"]["net_vat"]
    return {
        "season": s.as_dict(),
        "sales": out["sales"], "purchases": out["purchases"],
        "output_vat": out["sales"]["net_vat"], "input_vat": out["purchases"]["net_vat"],
        "payable": max(0, payable), "credit_carried_forward": max(0, -payable), "net": payable,
        "other_currency_invoices": left_out,
        "caveat": "Figures from your recorded invoices; confirm them in the tax organisation's portal before filing.",
    }


def reconciliation(db: Session, s: Season) -> dict[str, Any]:
    """The return and the TTMS report must describe the same trade: TTMS
    (all counterparties, مودیان sales included) = the return, side by side."""
    full = quarterly_report(db, s, include_moadian=True)
    ret = vat_return(db, s)
    rows = []
    for side, key in (("sales", "sales"), ("purchases", "purchases")):
        t = full["totals"][key]
        r = ret[side]
        rows.append({
            "side": side,
            "ttms_base": t["net_base"], "return_base": r["taxable_base"] + r["exempt_base"] - r["returns_base"],
            "ttms_vat": t["net_vat"], "return_vat": r["net_vat"],
        })
    for r in rows:
        r["matches"] = r["ttms_base"] == r["return_base"] and r["ttms_vat"] == r["return_vat"]
    return {"season": s.as_dict(), "rows": rows, "matches": all(r["matches"] for r in rows)}


def seasons_overview(db: Session, *, today: date | None = None, count: int = 6) -> list[dict[str, Any]]:
    """The last ``count`` seasons with their deadlines and how much trade each had."""
    today = today or date.today()
    s = season_of(today)
    out = []
    for _ in range(count):
        invoices = _season_invoices(db, s)
        out.append({
            **s.as_dict(),
            "current": s == season_of(today),
            "sales_count": sum(1 for i in invoices if i.kind == "sales"),
            "purchase_count": sum(1 for i in invoices if i.kind == "purchase"),
            "ttms_days_left": (s.ttms_deadline - today).days,
            "vat_days_left": (s.vat_deadline - today).days,
        })
        s = s.previous()
    return out


# --- the Excel file ---------------------------------------------------------------------------

_PARTY_HEADERS = ["ردیف", "نوع شخص", "شناسه/کد ملی", "کد اقتصادی", "نام", "کد پستی", "استان", "شهر", "نشانی",
                  "تلفن", "تعداد صورتحساب", "مبلغ قبل از مالیات (ریال)", "مالیات و عوارض ارزش افزوده (ریال)",
                  "مبلغ کل (ریال)", "برگشتی بدون مالیات (ریال)", "مالیات برگشتی (ریال)", "خالص مبلغ (ریال)",
                  "خالص مالیات (ریال)", "اطلاعات ناقص"]


def export_xlsx(db: Session, s: Season, *, include_moadian: bool = False) -> bytes:
    """Workbook for entering the season into TTMS: خلاصه, فروش, خرید, the
    invoices behind them and the parties whose details must be completed."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    rep = quarterly_report(db, s, include_moadian=include_moadian)
    ret = vat_return(db, s)
    rec = reconciliation(db, s)
    wb = Workbook()
    head = Font(bold=True)
    fill = PatternFill("solid", fgColor="E6F4F1")

    def sheet(title, headers, rows):
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head, fill, Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(r)
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(10, min(40, len(str(h)) + 4))
        ws.freeze_panes = "A2"
        return ws

    summary = wb.active
    summary.title = "خلاصه"
    summary.sheet_view.rightToLeft = True
    sd = rep["season"]
    for label, value in [
        ("فصل", sd["name"]), ("از", sd["start_jalali"]), ("تا", sd["end_jalali"]),
        ("مهلت گزارش معاملات فصلی (TTMS)", sd["ttms_deadline_jalali"]),
        ("مهلت اظهارنامه ارزش افزوده", sd["vat_deadline_jalali"]),
        ("", ""),
        ("فروش — مبلغ خالص (ریال)", rep["totals"]["sales"]["net_base"]),
        ("فروش — مالیات خالص (ریال)", rep["totals"]["sales"]["net_vat"]),
        ("خرید — مبلغ خالص (ریال)", rep["totals"]["purchases"]["net_base"]),
        ("خرید — مالیات خالص (ریال)", rep["totals"]["purchases"]["net_vat"]),
        ("فروش‌های ثبت‌شده در سامانه مودیان (کنار گذاشته شده)", rep["excluded_moadian"]["count"]),
        ("", ""),
        ("اظهارنامه — مالیات فروش (ریال)", ret["output_vat"]),
        ("اظهارنامه — اعتبار مالیاتی خرید (ریال)", ret["input_vat"]),
        ("اظهارنامه — قابل پرداخت (ریال)", ret["payable"]),
        ("اظهارنامه — بستانکار به دوره بعد (ریال)", ret["credit_carried_forward"]),
        ("تطبیق گزارش فصلی با اظهارنامه", "مطابق" if rec["matches"] else "مغایرت دارد"),
        ("طرف‌های معامله با اطلاعات ناقص", len(rep["incomplete"])),
        ("صورتحساب‌های ارزی (ریالی نشده)", len(rep["other_currency"])),
    ]:
        summary.append([label, value])
        summary.cell(summary.max_row, 1).font = head
    summary.column_dimensions["A"].width = 52
    summary.column_dimensions["B"].width = 22

    def party_rows(rows):
        return [[n, r["person_type"], r["national_id"], r["economic_code"], r["name"], r["postal_code"], r["province"],
                 r["city"], r["address"], r["phone"], r["invoice_count"], r["base"], r["vat"], r["total"],
                 r["returns_base"], r["returns_vat"], r["net_base"], r["net_vat"],
                 "، ".join(FIELD_LABELS[m] for m in r["missing"])] for n, r in enumerate(rows, start=1)]

    sheet("فروش", _PARTY_HEADERS, party_rows(rep["sales"]))
    sheet("خرید", _PARTY_HEADERS, party_rows(rep["purchases"]))
    sheet("صورتحساب‌ها", ["نوع", "شماره", "تاریخ", "طرف معامله", "ارز", "مبلغ قبل از مالیات", "مالیات", "مبلغ کل",
                          "وضعیت مودیان", "در گزارش"],
          [["فروش" if d["kind"] == "sales" else "خرید", d["number"], d["issue_date_jalali"], d["counterparty"] or "—",
            d["currency"], d["base"], d["vat"], d["total"], d["moadian_status"] or "—", "بله" if d["reported"] else "خیر"]
           for d in rep["details"]])
    sheet("اطلاعات ناقص", ["نوع", "طرف معامله", "کمبود"],
          [["فروش" if r["kind"] == "sales" else "خرید", r["name"], "، ".join(r["missing_labels"])]
           for r in rep["incomplete"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
