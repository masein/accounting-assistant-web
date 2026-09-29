"""The monthly close pack (roadmap §4.9, part 3)."""
from __future__ import annotations

import csv
import io
import uuid
import zipfile
from datetime import date

import pytest
from sqlalchemy import select

from app.db.tenant import use_company
from tests.test_statement_export import _company, _pdf_text, _rows, _xlsx

URL = "/manager-reports/close-pack"


@pytest.fixture()
def company(client, db):
    made = []

    def make(locale, currency, journals=()):
        api, cid = _company(client, db, locale, currency)
        made.append(cid)
        for d, desc, dr, cr, amount in journals:
            _post(api, d, desc, dr, cr, amount, currency)
        return api, cid
    yield make
    from tests.test_admin_audit import _purge_company
    for cid in made:
        _purge_company(db, cid)


def _post(api, d, desc, dr, cr, amount, currency):
    r = api.post("/transactions", json={"date": d, "description": desc, "currency": currency, "lines": [
        {"account_code": dr, "debit": amount, "credit": 0}, {"account_code": cr, "debit": 0, "credit": amount}]})
    assert r.status_code == 201, r.text
    return r.json()


def _checklist(api, month, **kw):
    r = api.get(f"{URL}/checklist", params={"month": month, **kw})
    assert r.status_code == 200, r.text
    return r.json()


def _states(data):
    return {i["key"]: i["state"] for i in data["items"]}


def _detail(data, key):
    return next(i["detail"] for i in data["items"] if i["key"] == key)


# --- an Iranian company: the whole pack, in Persian, by Jalali month -----------------------------------------------

IR = [("2026-08-01", "capital", "1110", "3110", 900_000_000),       # before Shahrivar: the opening balance
      ("2026-09-01", "sale", "1112", "4110", 400_000_000),          # 10 Shahrivar 1405
      ("2026-09-10", "rent", "6112", "1110", 90_000_000)]


def test_an_iranian_company_gets_its_pack_for_a_jalali_month(company, db):
    api, cid = company("ir", "IRR", IR)
    r = api.get(URL, params={"month": "1405-06"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/zip"
    assert 'filename="close-pack-1405-06.zip"' in r.headers["content-disposition"]
    z = zipfile.ZipFile(io.BytesIO(r.content))
    assert sorted(z.namelist()) == ["close-pack-1405-06.pdf", "close-pack-1405-06.xlsx", "journal-1405-06.csv"]

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(z.read("close-pack-1405-06.xlsx")))
    assert wb.sheetnames == ["فهرست بررسی", "صورت وضعیت مالی", "صورت سود و زیان", "صورت سود و زیان جامع",
                             "صورت تغییرات در حقوق مالکانه", "صورت جریان‌های نقدی", "تراز آزمایشی",
                             "سنی مطالبات", "سنی بدهی‌ها", "بودجه و عملکرد"]      # no bank lines, no bank sheets
    check = wb["فهرست بررسی"]
    assert check.sheet_view.rightToLeft is True
    assert [c.value for c in check[1]] == ["وضعیت", "مورد", "توضیح"]
    assert check["A2"].value == "انجام شده" and check["B2"].value == "برابری بدهکار و بستانکار"

    # the trial balance: opening from before the month, the month's movements, the closing balance
    tb = _rows(wb["تراز آزمایشی"])
    cash = next(v for k, v in tb.items() if k.startswith("1110"))
    assert cash == [900_000_000, 0, 90_000_000, 810_000_000]
    assert tb["جمع"] == [0, 490_000_000, 490_000_000, 0]
    assert wb["تراز آزمایشی"]["A3"].value.startswith("شهریور ۱۴۰۵")

    # the statements are the month's: Shahrivar is 23 August to 22 September 2026
    single = _xlsx(api.get("/manager-reports/financial/export", params={
        "from_date": "2026-08-23", "to_date": "2026-09-22", "format": "xlsx", "statements": "income_statement"}))
    assert _rows(wb["صورت سود و زیان"]) == _rows(single["صورت سود و زیان"])

    # every journal line of the month, with the Jalali date beside the ISO one
    text = z.read("journal-1405-06.csv").decode("utf-8-sig")
    lines = list(csv.DictReader(io.StringIO(text)))
    assert len(lines) == 4 and {ln["description"] for ln in lines} == {"sale", "rent"}
    first = lines[0]
    assert (first["date"], first["date_jalali"], first["account_code"], first["debit"]) == (
        "2026-09-01", "1405/06/10", "1112", "400000000")

    pages, _ = _pdf_text(api.get(URL, params={"month": "1405-06", "format": "pdf"}))
    assert pages >= 10                                                  # cover, 5 statements, 4 tables


def test_the_checklist_in_persian_or_english(company):
    api, _ = company("ir", "IRR", IR)
    fa = _checklist(api, "1405-06")
    assert fa["lang"] == "fa" and fa["label"] == "شهریور ۱۴۰۵"
    assert (fa["from_date"], fa["to_date"]) == ("2026-08-23", "2026-09-22")
    assert fa["items"][0]["item"] == "برابری بدهکار و بستانکار" and "۴۹۰,۰۰۰,۰۰۰" in fa["items"][0]["detail"]
    en = _checklist(api, "1405-06", lang="en")
    assert en["label"] == "Shahrivar 1405" and en["items"][0]["item"] == "Debits equal credits"
    assert "490,000,000" in en["items"][0]["detail"]


# --- a UK company: what the checklist catches ----------------------------------------------------------------------

UK = [("2026-07-01", "capital", "1200", "3000", 90_000), ("2026-08-10", "sale", "1100", "4000", 40_000),
      ("2026-08-15", "rent", "7100", "1200", 9_000)]


def _statement(db, cid, rows):
    from app.models.bank_statement import BankStatement, BankStatementRow
    with use_company(cid):
        s = BankStatement(bank_name="Barclays", account_number="12345678", source_type="csv",
                          source_filename="aug.csv", currency="GBP", status="reviewing")
        db.add(s)
        db.flush()
        for i, (d, desc, debit, credit, balance, status) in enumerate(rows):
            db.add(BankStatementRow(statement_id=s.id, row_index=i, tx_date=d, description=desc, debit=debit,
                                    credit=credit, balance=balance, recon_status=status))
        db.commit()
        return s.id


def test_a_uk_company_reads_the_checklist_in_persian_but_its_pack_stays_english(company):
    api, _ = company("uk", "GBP", UK)
    fa = _checklist(api, "2026-08", lang="fa")
    assert fa["lang"] == "fa" and fa["items"][0]["item"] == "برابری بدهکار و بستانکار"
    assert "۴۹,۰۰۰" in fa["items"][0]["detail"]
    assert "۲۰۲۶/۰۸/۳۱" in _detail(fa, "lock")                               # not "2026-08-31" inside Persian
    wb = _xlsx(api.get(URL, params={"month": "2026-08", "format": "xlsx", "lang": "fa"}))
    assert wb.sheetnames[0] == "Checklist" and wb.worksheets[0].sheet_view.rightToLeft is not True


def test_a_clean_month_passes(company):
    api, _ = company("uk", "GBP", UK)
    data = _checklist(api, "2026-08")
    assert [i["key"] for i in data["items"]] == ["balanced", "bank", "drafts", "fx", "lock"]   # no payroll, no assets
    assert _states(data) == {"balanced": "ok", "bank": "info", "drafts": "ok", "fx": "ok", "lock": "info"}
    assert data["open"] == 0 and data["summary"] == "Everything on the checklist is done."
    assert _detail(data, "balanced") == "4 journal lines this month, 49,000 each side."


def test_the_checklist_catches_what_is_still_open(company, db):
    from app.models.employee_pay import EmployeePayProfile
    api, cid = company("uk", "GBP", UK)
    _statement(db, cid, [(date(2026, 8, 10), "Card payment", 0, 40_000, 121_000, "matched"),
                         (date(2026, 8, 20), "Unknown transfer", 250, 0, 120_750, "unmatched"),
                         (date(2026, 8, 20), "Unknown transfer", 250, 0, 120_750, "duplicate")])
    client = api.post("/entities", json={"type": "client", "name": "Aria Trading"}).json()
    r = api.post("/invoices", json={"number": "D-1", "kind": "sales", "status": "draft", "issue_date": "2026-08-20",
                                    "due_date": "2026-09-20", "amount": 1_000, "currency": "GBP",
                                    "entity_id": client["id"]})
    assert r.status_code == 201, r.text
    emp = api.post("/entities", json={"type": "employee", "name": "Sara Ahmadi"}).json()
    with use_company(cid):
        db.add(EmployeePayProfile(entity_id=uuid.UUID(emp["id"]), base_salary=3_000, currency="GBP"))
        db.commit()
    _post(api, "2026-08-12", "a dollar bill", "7100", "1200", 100, "USD")    # no USD rate yet

    data = _checklist(api, "2026-08")
    assert _states(data) == {"balanced": "ok", "bank": "warn", "drafts": "warn", "payroll": "warn", "fx": "warn",
                             "lock": "info"}
    assert _detail(data, "bank") == "1 of 3 bank lines this month aren't matched yet."
    assert _detail(data, "drafts") == "Invoices dated this month still in draft: 1 (D-1)."
    assert _detail(data, "payroll") == "No pay run ends this month; employees on payroll: 1."
    assert "(USD)" in _detail(data, "fx")
    assert data["open"] == 4 and data["summary"] == "Checklist items that still need attention: 4."

    # a draft pay run is still open; posted, it's done
    from app.models.pay_run import PayRun
    with use_company(cid):
        run = PayRun(period_start=date(2026, 8, 1), period_end=date(2026, 8, 31), pay_date=date(2026, 8, 31),
                     currency="GBP", status="draft")
        db.add(run)
        db.commit()
        run_id = run.id
    assert _detail(_checklist(api, "2026-08"), "payroll") == "A pay run ending this month is still a draft."
    with use_company(cid):
        db.execute(select(PayRun).where(PayRun.id == run_id)).scalar_one().status = "posted"
        from app.services.period_service import set_closed_period
        set_closed_period(db, date(2026, 8, 31))
        db.commit()
    data = _checklist(api, "2026-08")
    assert _states(data)["payroll"] == "ok" and _states(data)["lock"] == "ok"
    assert _detail(data, "lock").startswith("The books are locked through")
    # the other months don't see August's loose ends
    assert _states(_checklist(api, "2026-07"))["drafts"] == "ok"


def test_depreciation_not_run_is_caught(company, db):
    from app.services import fixed_assets as fa
    api, cid = company("uk", "GBP", UK)
    r = api.post("/fixed-assets", json={"name": "Desks", "category": "furniture", "cost": 3_600, "life_months": 36,
                                        "acquired_on": "2026-05-04", "depreciation_start": "2026-05-01"})
    assert r.status_code in (200, 201), r.text
    data = _checklist(api, "2026-08")
    assert _states(data)["depreciation"] == "warn" and "not posted yet: 4" in _detail(data, "depreciation")
    with use_company(cid):
        fa.run(db, date(2026, 8, 31))
        db.commit()
    assert _states(_checklist(api, "2026-08"))["depreciation"] == "ok"


def test_the_pack_tables_aging_bank_and_budgets(company, db):
    from app.models.account import Account
    from app.models.budget import BudgetLimit
    api, cid = company("uk", "GBP", UK)
    a = api.post("/entities", json={"type": "client", "name": "Aria Trading"}).json()
    b = api.post("/entities", json={"type": "client", "name": "Borna Co"}).json()
    s = api.post("/entities", json={"type": "supplier", "name": "Paper Ltd"}).json()
    for number, kind, ent, issued, due, amount in (
            ("S-1", "sales", a, "2026-06-01", "2026-07-01", 10_000),       # 61 days late at 31 August
            ("S-2", "sales", b, "2026-08-20", "2026-09-20", 4_000),        # not due yet
            ("S-3", "sales", b, "2026-09-02", "2026-09-30", 7_000),        # after the month: not in it
            ("P-1", "purchase", s, "2026-08-05", "2026-08-10", 2_000)):    # 21 days late
        r = api.post("/invoices", json={"number": number, "kind": kind, "status": "issued", "issue_date": issued,
                                        "due_date": due, "amount": amount, "currency": "GBP", "entity_id": ent["id"]})
        assert r.status_code == 201, r.text
    _statement(db, cid, [(date(2026, 7, 30), "July", 0, 90_000, 90_000, "unmatched"),       # not this month
                         (date(2026, 8, 10), "Card payment", 0, 40_000, 130_000, "matched"),
                         (date(2026, 8, 20), "Unknown transfer", 250, 0, 129_750, "unmatched"),
                         (date(2026, 8, 21), "Skipped fee", 5, 0, 129_745, "skipped")])
    with use_company(cid):
        rent = db.execute(select(Account.name).where(Account.code == "7100")).scalar_one()
        db.add(BudgetLimit(month="2026-08", category=rent, limit_amount=12_000))
        db.commit()

    r = api.get(URL, params={"month": "2026-08", "format": "xlsx"})
    wb = _xlsx(r)
    assert 'filename="close-pack-2026-08.xlsx"' in r.headers["content-disposition"]
    assert wb.sheetnames == ["Checklist", "Balance sheet", "Profit and loss account", "Comprehensive income",
                             "Changes in equity", "Cash flows", "Trial balance", "Receivables aging", "Payables aging",
                             "Bank reconciliation", "Bank lines not matched", "Budget vs actual"]
    ar = _rows(wb["Receivables aging"])
    assert ar["Aria Trading"] == [0, 0, 0, 10_000, 0, 10_000]
    assert ar["Borna Co"] == [4_000, 0, 0, 0, 0, 4_000]
    assert ar["Total"] == [4_000, 0, 0, 10_000, 0, 14_000]
    assert list(ar)[:2] == ["Aria Trading", "Borna Co"]                        # the biggest balance first
    assert _rows(wb["Payables aging"])["Paper Ltd"] == [0, 2_000, 0, 0, 0, 2_000]
    assert _rows(wb["Bank reconciliation"])["Barclays · 12345678 · GBP"] == [3, 1, 1, 1, 129_745]
    unmatched = _rows(wb["Bank lines not matched"])
    assert list(unmatched) == ["Unknown transfer"] and unmatched["Unknown transfer"][1:] == [250, None]
    budget = _rows(wb["Budget vs actual"])
    assert budget[rent] == [12_000, 9_000, 3_000, "75%"] and budget["Total"] == [12_000, 9_000, 3_000, "75%"]
    cells = [c for row in wb["Budget vs actual"].iter_rows(min_row=6) for c in row[1:4] if c.value is not None]
    assert all(isinstance(c.value, int) and c.number_format == "#,##0;(#,##0);-" for c in cells)

    r = api.get(URL, params={"month": "2026-08", "format": "pdf"})
    from pypdf import PdfReader
    texts = [p.extract_text() or "" for p in PdfReader(io.BytesIO(r.content)).pages]
    aging_page = next(t for t in texts if "Receivables aging" in t)
    assert "Payables aging" in aging_page                                        # short tables share a page
    bank_page = next(t for t in texts if "Bank reconciliation" in t)
    assert "Bank lines not matched" in bank_page and "Budget vs actual" in bank_page
    assert aging_page.count("Export Co uk") == 1                                 # one letterhead a page
    pages, text = _pdf_text(r)
    assert "Monthly close pack" in text and "Aug 2026" in text and "Receivables aging" in text
    assert "Bank reconciliation" in text and "Checklist items that still need attention: 1." in text
    assert "Unknown transfer" in text


def test_an_empty_month_still_makes_a_pack(company):
    api, _ = company("uk", "GBP")
    wb = _xlsx(api.get(URL, params={"month": "2026-08", "format": "xlsx"}))
    assert "Bank reconciliation" not in wb.sheetnames
    assert wb["Receivables aging"]["A7"].value.startswith("Nothing open at")          # the note, no rows
    assert wb["Budget vs actual"]["A7"].value == "No budgets are set for this month."
    assert _rows(wb["Trial balance"]) == {"Total": [0, 0, 0, 0],
                                          "Balances are debit positive, credit in brackets.": [None] * 4}


def test_the_default_month_and_the_picker(company):
    from app.services.calendar_periods import GREGORIAN, JALALI, month_key, previous_month_key
    api, _ = company("uk", "GBP", UK)
    data = api.get(f"{URL}/checklist").json()
    assert data["month"] == previous_month_key(date.today(), GREGORIAN)            # last month, by default
    assert len(data["months"]) == 12 and data["months"][0]["key"] == month_key(date.today(), GREGORIAN)
    ir, _ = company("ir", "IRR", IR)
    data = ir.get(f"{URL}/checklist").json()
    assert data["month"] == previous_month_key(date.today(), JALALI)               # an Iranian company: Jalali
    assert int(data["months"][0]["key"][:4]) < 1700


def test_what_the_pack_refuses(company):
    api, _ = company("uk", "GBP", UK)
    for bad in ("2026-13", "2026-8", "august"):
        assert api.get(URL, params={"month": bad}).status_code == 422, bad
        assert api.get(f"{URL}/checklist", params={"month": bad}).status_code == 422, bad
    assert api.get(URL, params={"month": "2026-08", "format": "csv"}).status_code == 422
    assert api.get(URL, params={"month": "2026-08", "lang": "de"}).status_code == 422


def test_who_may_download_the_pack():
    from app.core.auth import SessionUser
    from app.core.permissions import user_can_access
    for role, ok in (("owner", True), ("cfo", True), ("accountant", True), ("viewer", True), ("manager", False),
                     ("employee", False), ("personal", True)):              # a personal user's own books
        u = SessionUser(user_id="u", username="u", role=role, is_admin=False)
        for path in (URL, f"{URL}/checklist"):
            assert user_can_access(u, "GET", path) is ok, (role, path)


def test_the_close_pack_panel_is_wired():
    from tests.i18n_source import i18n_text
    html = open("app/static/index.html", encoding="utf-8").read()
    for el in ('id="close-pack-card"', 'id="close-pack-month"', 'id="close-pack-zip"', 'id="close-pack-checklist"'):
        assert el in html, el
    assert html.count('close-pack-file"') == 2
    js = open("app/static/js/05-reports-manager.js", encoding="utf-8").read()
    assert "'/manager-reports/close-pack/checklist?'" in js and "'/manager-reports/close-pack?'" in js
    ops = open("app/static/js/12-ops.js", encoding="utf-8").read()
    assert "loadClosePack();" in ops
    text = i18n_text()
    for k in ("closePackTitle", "closePackHint", "closePackMonth", "closePackZip", "closePackPdf", "closePackXlsx",
              "closePackFailed"):
        assert text.count(f"{k}:") == 4, k
