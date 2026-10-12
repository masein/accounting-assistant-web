"""Pence in what the app writes (docs/PLAN_PENCE.md, phase P5; scenarios Q11a–e).

The ledger kept £7.50 from P2 on, but the sentences, averages, files and
cheques still cut it: the accountant said «£23,915» for £23,915.50, a burn rate
of £1,000.02 was «£1,000», a £9.99 subscription was «£9», a CSV cell read 7.5,
Excel showed «8», and a £1,234.56 cheque printed «£1,234.00». Whole amounts
and rial text read exactly as before."""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

import pytest

from app.core.money import (
    average, grouped, grouped_fa, persian_figures, sheet_format, to_places,
)
from app.db.tenant import use_company
from tests.test_statement_export import _company


@pytest.fixture()
def uk(client, db):
    from tests.test_admin_audit import _purge_company
    api, cid = _company(client, db, "uk", "GBP")
    yield api, cid
    client.cookies.clear()
    _purge_company(db, cid)


@pytest.fixture()
def no_model(monkeypatch):
    from app.api import ai_accountant as api_mod

    async def refuse(*a, **kw):
        raise AssertionError("the model was called for a question the books answer")
    monkeypatch.setattr(api_mod, "run_chat_turn", refuse)


def _post(api, when, desc, debit_code, credit_code, amount):
    r = api.post("/transactions", json={"date": str(when), "description": desc, "lines": [
        {"account_code": debit_code, "debit": amount, "credit": 0},
        {"account_code": credit_code, "debit": 0, "credit": amount}]})
    assert r.status_code == 201, r.text


def _month_starts(n: int) -> list[date]:
    """The first day of this month and the n - 1 before it, oldest first."""
    y, m = date.today().year, date.today().month
    out = []
    for _ in range(n):
        out.append(date(y, m, 1))
        y, m = (y - 1, 12) if m == 1 else (y, m - 1)
    return out[::-1]


# --- the helpers --------------------------------------------------------------------------------------

def test_an_amount_in_a_sentence():
    assert grouped(Decimal("23915.5")) == "23,915.50"            # it read «23,915.5»
    assert grouped(Decimal("-7.5")) == "-7.50"
    assert grouped(Decimal("0.125")) == "0.125"                  # a unit price keeps its four places
    assert grouped(950) == "950" and grouped(Decimal("950.00")) == "950"
    assert grouped(2_500_000) == "2,500,000"                     # rial text doesn't move
    assert grouped(None) == "" and grouped("n/a") == "n/a"


def test_figures_in_persian():
    assert grouped_fa(2_500_000) == "۲٬۵۰۰٬۰۰۰"
    assert grouped_fa(Decimal("7.5")) == "۷٫۵۰"                  # «۷.۵۰» read as a stray full stop
    # a full stop that ends the sentence stays one
    assert persian_figures("مبلغ 1,250.75 پوند.") == "مبلغ ۱٬۲۵۰٫۷۵ پوند."
    assert persian_figures("1,234.5", thousands=",") == "۱,۲۳۴٫۵"   # pages that always kept the comma


def test_q11b_averages_to_the_penny():
    assert average([Decimal("1000.10"), Decimal("999.95"), 1000], "GBP") == Decimal("1000.02")
    assert average([1000, 1001, 1001], "IRR") == 1000             # whole rials, truncated as before
    assert average([], "GBP") == 0
    assert to_places(9.99, "GBP") == Decimal("9.99") and to_places(9.99, "IRR") == 9
    assert sheet_format("GBP") == "#,##0.00" and sheet_format("IRR") == "#,##0" and sheet_format("KWD") == "#,##0.000"


# --- Q11a: the accountant's sentences -----------------------------------------------------------------

def test_q11a_the_quick_answer_keeps_its_pence(uk, no_model):
    api, _ = uk
    _post(api, date.today(), "Capital", "1200", "3000", 24_000)
    _post(api, date.today(), "Monthly account fee", "8000", "1200", "7.50")
    r = api.post("/ai-accountant/chat", json={"message": "how much cash do we have?"})
    assert r.status_code == 200, r.text
    assert r.json()["text"] == "Cash and bank today: £23,992.50."      # it said «£23,992»
    _post(api, date.today(), "Rent", "8000", "1200", "0.50")
    assert api.post("/ai-accountant/chat", json={"message": "how much cash do we have?"}).json()["text"] \
        == "Cash and bank today: £23,992."                              # a whole amount as before


def test_q11a_the_cfo_report_writes_pence():
    from app.services import cfo_intelligence as cfo
    assert cfo._s("en", "qa_burn").format(burn=Decimal("1000.02")) == "Burn rate: 1,000.02/month."
    assert cfo._s("en", "qa_burn").format(burn=1_250_000) == "Burn rate: 1,250,000/month."
    assert cfo._s("en", "ins_runway_critical_title").format(runway=2.25) == "Cash runway critical: 2.2 months"


def test_q11a_an_insight_writes_the_base_currency_places():
    from app.services import insight_service as ins
    token = ins._BASE.set("GBP")
    try:
        assert ins._fmt(3000.05 / 3) == "1,000.02" and ins._fmt(950) == "950"
    finally:
        ins._BASE.reset(token)
    token = ins._BASE.set("IRR")
    try:
        assert ins._fmt(2_500_000.6) == "2,500,001"                    # rials round as they always did
    finally:
        ins._BASE.reset(token)


# --- Q11b: averages ------------------------------------------------------------------------------------

def test_q11b_the_burn_rate_is_to_the_penny(uk):
    api, _ = uk
    first, second, third = _month_starts(3)
    _post(api, first, "Capital", "1200", "3000", 24_000)
    for when, amount in ((first, "1000.10"), (second, "999.95"), (third, "1000.00")):
        _post(api, when, "Supplies", "8000", "1200", amount)
    dash = api.get("/reports/owner-dashboard")
    assert dash.status_code == 200, dash.text
    kpi = {k["key"]: k for k in dash.json()["kpis"]}
    assert kpi["burn_rate"]["value"] == 1000.02                         # it was 1000
    cfo = api.get("/brain/cfo/report")
    assert cfo.status_code == 200, cfo.text
    assert cfo.json()["burn_rate"] == 1000.02
    assert {k["key"]: k["value"] for k in cfo.json()["kpis"]}["burn_rate"] == 1000.02


def test_q11b_a_subscription_and_a_goal_keep_their_pence(uk, db):
    api, cid = uk
    for when in _month_starts(4):
        _post(api, when, "Spotify", "8000", "1200", "9.99")
    _post(api, date.today(), "To savings", "1210", "1200", 100)
    from app.services.personal_goals import progress
    from app.services.recurring_detection import detect_recurring
    with use_company(cid):
        spotify = next(c for c in detect_recurring(db) if "spotify" in c.description.lower())
        assert spotify.typical_amount == Decimal("9.99")              # it was 9
        goal = SimpleNamespace(id="g1", name="Holiday", account_code="1210", target_amount=1000,
                               target_date=None, archived=False)
        row = progress(db, [goal])[0]
    assert row["current"] == 100 and row["pace_per_month"] == Decimal("33.33")


def test_q11b_what_a_goal_needs_each_month_is_rounded_up_to_the_penny():
    from app.services.personal_goals import _per_month
    assert _per_month(100, 3, "GBP") == Decimal("33.34")              # it was £34
    assert _per_month(1_000_000, 3, "IRR") == 333_334


# --- Q11c: notifications and digests -------------------------------------------------------------------

def test_q11c_the_digest_writes_pence():
    from app.services.digest_service import format_digest
    body = format_digest("Acme", {"currency": "GBP", "low_cash": False, "low_cash_reasons": [], "runway_months": 4.2,
                                  "cash_on_hand": Decimal("1250.75"), "ar_outstanding": Decimal("30.5"),
                                  "ar_overdue": 0, "ap_outstanding": 950, "ap_overdue": 0})
    assert "1,250.75" in body and "30.50" in body and "950" in body and "950.00" not in body


# --- Q11d: files ---------------------------------------------------------------------------------------

def test_q11d_csv_and_excel_keep_pence(uk):
    from openpyxl import load_workbook
    api, _ = uk
    _post(api, date.today(), "Monthly account fee", "8000", "1200", "7.50")
    _post(api, date.today(), "Rent", "8000", "1200", 950)
    csv = api.get("/exports/transactions.csv").content.decode("utf-8-sig")
    assert ",7.50," in csv and ",950," in csv and ",7.5," not in csv
    wb = load_workbook(io.BytesIO(api.get("/exports/transactions.xlsx").content))
    cells = [(row[6].value, row[6].number_format) for row in wb.active.iter_rows(min_row=2) if row[6].value]
    assert (7.5, "#,##0.00") in cells and (950, "#,##0.00") in cells
    year = date.today().year
    r = api.get("/manager-reports/financial/export",
                params={"from_date": f"{year}-01-01", "to_date": f"{year}-12-31", "format": "xlsx"})
    assert r.status_code == 200, r.text
    formats = {c.number_format for ws in load_workbook(io.BytesIO(r.content)) for row in ws.iter_rows()
               for c in row if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)}
    assert "#,##0.00;(#,##0.00);-" in formats and "#,##0;(#,##0);-" not in formats


def test_q11d_a_csv_cell_with_pence():
    from app.core.spreadsheet import csv_safe
    assert csv_safe(Decimal("7.5")) == "7.50" and csv_safe(Decimal("950")) == 950
    assert csv_safe(Decimal("0.125")) == "0.125" and csv_safe(7.5) == 7.5


# --- Q11e: cheques -------------------------------------------------------------------------------------

def test_q11e_a_cheque_to_the_penny(uk, db):
    from app.services import cheque_print as cp
    _, cid = uk
    row = SimpleNamespace(amount=Decimal("1234.56"), due_date=date(2026, 10, 12), counterparty="Acme Ltd",
                          entity_id=None)
    with use_company(cid):
        v = cp.cheque_values(db, row)
    assert v["amount"] == "£1,234.56"                                    # it printed «£1,234.00»
    assert v["amount_words"] == "One thousand, two hundred and thirty-four pounds and fifty-six pence only"
