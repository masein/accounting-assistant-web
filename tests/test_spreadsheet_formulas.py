"""Text in an export never runs as a spreadsheet formula (CSV/XLSX formula
injection). Journal descriptions, references and party names come from users,
bank statements, SMS and imports; one that starts with = + - @ used to reach
the manager-report CSVs and the close-pack journal as it was, and openpyxl
stored an "=…" description as a live formula in every workbook."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import Path

import pytest

from app.core.spreadsheet import csv_safe, no_formulas
from tests.test_statement_export import _company

EVIL = '=HYPERLINK("http://evil.example/x","Pay here")'
AT = "@SUM(1+1)*cmd|' /C calc'!A0"
PLUS = "+1+2"


@pytest.mark.parametrize("value, out", [
    (EVIL, "'" + EVIL), (AT, "'" + AT), (PLUS, "'+1+2"), ("-2+3", "'-2+3"), ("\t=1", "'\t=1"), ("\r=1", "'\r=1"),
    ("-1500", "-1500"), ("1,250.50", "1,250.50"), ("+98912", "+98912"), ("hello = world", "hello = world"),
    ("", ""), (None, ""), (42, 42), (-7.5, -7.5), ("'=already", "'=already"),
])
def test_csv_cells(value, out):
    assert csv_safe(value) == out
    assert csv_safe(csv_safe(value)) == out                      # applying it twice changes nothing


def _load(body: bytes):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(body))


def _formulas(body: bytes) -> list[str]:
    return [f"{ws.title}!{c.coordinate}" for ws in _load(body).worksheets
            for row in ws.iter_rows() for c in row if c.data_type == "f"]


def _texts(body: bytes) -> list[str]:
    return [c.value for ws in _load(body).worksheets for row in ws.iter_rows() for c in row
            if isinstance(c.value, str)]


def test_a_workbook_keeps_the_text_but_not_the_formula():
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append([EVIL, 5, "plain"])
    assert wb.active["A1"].data_type == "f"                      # what openpyxl does on its own
    buf = io.BytesIO()
    no_formulas(wb).save(buf)
    ws = _load(buf.getvalue()).active
    assert ws["A1"].data_type == "s" and ws["A1"].value == EVIL   # the text, exactly as typed
    assert ws["B1"].value == 5 and ws["C1"].value == "plain"


@pytest.fixture()
def books(client, db):
    api, cid = _company(client, db, "uk", "GBP")
    r = api.post("/transactions", json={"date": "2026-08-15", "reference": PLUS, "description": EVIL, "currency": "GBP",
                                        "lines": [{"account_code": "7100", "debit": 9_000, "credit": 0, "line_description": AT},
                                                  {"account_code": "1200", "debit": 0, "credit": 9_000}]})
    assert r.status_code == 201, r.text
    yield api, cid
    from tests.test_admin_audit import _purge_company
    client.cookies.clear()
    _purge_company(db, cid)


def _csv(body: bytes | str) -> list[dict]:
    text = body.decode("utf-8-sig") if isinstance(body, bytes) else body
    return list(csv.DictReader(io.StringIO(text)))


def test_the_journal_exports(books):
    api, _ = books
    # transactions CSV and workbook
    rows = [r for r in _csv(api.get("/exports/transactions.csv").content) if "HYPERLINK" in r["description"]]
    assert rows and all(r["description"] == "'" + EVIL and r["reference"] == "'+1+2" for r in rows)
    assert "'" + AT in {r["line_description"] for r in rows}
    assert {r["debit"] for r in rows} <= {"9000", "9000.00", "0", "0.00"}          # numbers are left alone
    body = api.get("/exports/transactions.xlsx").content
    assert _formulas(body) == [] and EVIL in _texts(body) and AT in _texts(body)   # no apostrophe needed there
    # the manager reports' general journal and general ledger
    aug = {"format": "csv", "from_date": "2026-08-01", "to_date": "2026-08-31"}
    rows = _csv(api.get("/manager-reports/books/general-journal", params=aug).text)
    assert "'" + EVIL in {r["description"] for r in rows} and "'" + AT in {r["line_description"] for r in rows}
    ledger = api.get("/manager-reports/books/general-ledger", params=aug).text
    assert ledger.count("\n") > 1 and not re.search(r'(^|,)"?[=@+]', ledger, re.M), ledger[:400]
    # the close pack: the journal CSV and the workbook inside the ZIP
    r = api.get("/manager-reports/close-pack", params={"month": "2026-08", "format": "zip"})
    assert r.status_code == 200, r.text
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        journal = _csv(z.read("journal-2026-08.csv"))
        workbook = z.read("close-pack-2026-08.xlsx")
    assert "'" + EVIL in {j["description"] for j in journal} and "'+1+2" in {j["reference"] for j in journal}
    assert _formulas(workbook) == []


def test_party_names_in_the_ttms_workbook(client, db):
    from tests.test_admin_audit import _purge_company
    api, cid = _company(client, db, "ir", "IRR")
    try:
        buyer = api.post("/entities", json={"type": "client", "name": EVIL, "national_id": "0012345678",
                                            "economic_code": "411111111111"}).json()
        r = api.post("/invoices", json={"number": "F-1", "kind": "sales", "status": "issued", "issue_date": "2026-07-01",
                                        "due_date": "2026-07-30", "amount": 5_000_000, "currency": "IRR",
                                        "entity_id": buyer["id"]})
        assert r.status_code == 201, r.text
        r = api.get("/tax/ir/quarterly/export", params={"year": 1405, "season": 2})
        assert r.status_code == 200, r.text
        assert _formulas(r.content) == [] and EVIL in _texts(r.content)
    finally:
        client.cookies.clear()
        _purge_company(db, cid)


def test_every_writer_uses_the_guard():
    """New exports get it too: no bare csv.writer, and every workbook the app
    builds is passed through no_formulas before it's saved."""
    app = Path(__file__).resolve().parents[1] / "app"
    bare, unguarded = [], []
    for p in app.rglob("*.py"):
        if p.name == "spreadsheet.py":
            continue
        src = p.read_text(encoding="utf-8")
        if re.search(r"\bcsv\.(writer|DictWriter)\(", src):
            bare.append(str(p.relative_to(app)))
        if "Workbook()" in src and "no_formulas(" not in src:
            unguarded.append(str(p.relative_to(app)))
    assert bare == [] and unguarded == [], (bare, unguarded)
