"""The financial statements as PDF and Excel, made on the server (roadmap §4.9)."""
from __future__ import annotations

import io
import uuid

import pytest


def _company(client, db, locale, currency):
    from app.core.auth import CSRF_COOKIE, create_session_token, generate_csrf_token
    from app.core.config import settings
    from app.db.seed import seed_chart_if_empty
    from app.db.tenant import use_company
    from app.models.company import Company
    from tests.conftest import _CSRFTestClient
    c = Company(id=uuid.uuid4(), name=f"Export Co {locale}", slug=f"exp-{uuid.uuid4().hex[:6]}", locale=locale,
                base_currency=currency, status="active", token_version=0)
    db.add(c)
    db.commit()
    cid = str(c.id)
    with use_company(cid):
        seed_chart_if_empty(db, locale=locale if locale in ("ir", "uk") else "ir")
        if currency != "IRR":
            from app.services.fx_service import set_reporting_currency
            set_reporting_currency(db, currency)
        if locale not in ("ir", "uk"):
            from app.services.locale_service import set_reporting_locale
            set_reporting_locale(db, "default")
        db.commit()
    csrf = generate_csrf_token()
    client.cookies.clear()
    client.cookies.set(settings.auth_cookie_name, create_session_token(
        user_id=str(uuid.uuid4()), username="owner", is_admin=True, role="owner", company_id=cid))
    client.cookies.set(CSRF_COOKIE, csrf)
    db.expunge_all()
    return _CSRFTestClient(client, csrf), cid


@pytest.fixture()
def books(client, db):
    made = []

    def make(locale, currency, journals):
        api, cid = _company(client, db, locale, currency)
        made.append(cid)
        for d, desc, dr, cr, amount in journals:
            r = api.post("/transactions", json={"date": d, "description": desc, "currency": currency, "lines": [
                {"account_code": dr, "debit": amount, "credit": 0}, {"account_code": cr, "debit": 0, "credit": amount}]})
            assert r.status_code == 201, r.text
        return api
    yield make
    from tests.test_admin_audit import _purge_company
    for cid in made:
        _purge_company(db, cid)


IR = [("2026-04-01", "capital", "1110", "3110", 900_000_000), ("2026-05-10", "sale", "1112", "4110", 400_000_000),
      ("2026-06-15", "rent", "6112", "1110", 90_000_000)]
UK = [("2026-04-01", "capital", "1200", "3000", 90_000), ("2026-05-10", "sale", "1100", "4000", 40_000),
      ("2026-06-15", "rent", "7100", "1200", 9_000)]


def _xlsx(r):
    from openpyxl import load_workbook
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    return load_workbook(io.BytesIO(r.content))


def _pdf_text(r):
    from pypdf import PdfReader
    assert r.status_code == 200, r.text
    assert r.content[:4] == b"%PDF" and r.headers["content-type"] == "application/pdf"
    reader = PdfReader(io.BytesIO(r.content))
    return len(reader.pages), "\n".join(p.extract_text() or "" for p in reader.pages)


def _rows(ws):
    """label → list of numbers, from the table under the header row."""
    out = {}
    for row in ws.iter_rows(min_row=6, values_only=True):
        if row and row[0]:
            out[str(row[0]).strip()] = [v for v in row[1:]]
    return out


P = {"from_date": "2026-01-01", "to_date": "2026-09-30"}


def test_an_iranian_company_gets_its_five_statements_in_persian(books):
    api = books("ir", "IRR", IR)
    wb = _xlsx(api.get("/manager-reports/financial/export", params={**P, "format": "xlsx"}))
    assert wb.sheetnames == ["صورت وضعیت مالی", "صورت سود و زیان", "صورت سود و زیان جامع",
                             "صورت تغییرات در حقوق مالکانه", "صورت جریان‌های نقدی"]
    bs = wb["صورت وضعیت مالی"]
    assert bs.sheet_view.rightToLeft is True
    # three date columns on the Iranian balance sheet, in Jalali with Persian digits
    assert [c.value for c in bs[5]] == ["شرح", "۱۴۰۵/۰۷/۰۸", "۱۴۰۴/۰۷/۰۸", "۱۴۰۳/۰۷/۰۹"]
    # the workbook carries the same figures as the statement the page shows
    js = api.get("/manager-reports/financial/iran/income-statement", params=P).json()
    sheet = _rows(wb["صورت سود و زیان"])
    for r in js["rows"]:
        if r["row_type"] in ("header", "spacer") or r["amount_current"] is None:
            continue
        want = -abs(r["amount_current"]) if r["is_negative_presentation"] else r["amount_current"]
        assert sheet[r["label_fa"]][0] == want, r["label_fa"]
    # numbers are numbers, totals are bold
    cells = [c for row in wb["صورت سود و زیان"].iter_rows(min_row=6) for c in row[1:2] if c.value is not None]
    assert all(isinstance(c.value, int) and c.number_format == "#,##0;(#,##0);-" for c in cells)
    assert any(row[0].font.bold for row in wb["صورت سود و زیان"].iter_rows(min_row=6) if row[0].value)
    pages, _text = _pdf_text(api.get("/manager-reports/financial/export", params={**P, "format": "pdf"}))
    assert pages >= 5                                                   # one statement a page (at least)


def test_english_labels_for_an_iranian_company_on_request(books):
    api = books("ir", "IRR", IR)
    wb = _xlsx(api.get("/manager-reports/financial/export",
                       params={**P, "format": "xlsx", "lang": "en", "statements": "income_statement"}))
    assert wb.sheetnames == ["Income statement"] and wb.active.sheet_view.rightToLeft is not True


def test_a_uk_company_gets_frs102_statements(books):
    api = books("uk", "GBP", UK)
    wb = _xlsx(api.get("/manager-reports/financial/export", params={**P, "format": "xlsx", "lang": "fa"}))
    assert wb.sheetnames == ["Balance sheet", "Profit and loss account", "Comprehensive income",
                             "Changes in equity", "Cash flows"]           # no Persian UK template; Excel's 31 characters
    assert wb["Comprehensive income"]["A2"].value == "Statement of comprehensive income"
    assert "GBP" in wb["Balance sheet"]["A3"].value
    js = api.get("/manager-reports/financial/uk/balance-sheet", params={"as_of": P["to_date"]}).json()
    sheet = _rows(wb["Balance sheet"])
    for r in js["rows"]:
        if r["row_type"] not in ("header", "spacer") and r["amount_current"] is not None:
            want = -abs(r["amount_current"]) if r["is_negative_presentation"] else r["amount_current"]
            assert sheet[r["label"]][0] == want, r["label"]
    equity = wb["Changes in equity"]
    assert [c.value for c in equity[5]][-1] == "Total"                     # the matrix: components, then total
    pages, text = _pdf_text(api.get("/manager-reports/financial/export",
                                    params={**P, "format": "pdf", "statements": "balance_sheet,income_statement"}))
    assert pages >= 2 and "Balance sheet" in text and "Profit and loss account" in text and "Export Co uk" in text
    assert "Amounts in GBP" in text


def test_a_company_on_the_generic_statements(books):
    api = books("default", "IRR", IR)
    wb = _xlsx(api.get("/manager-reports/financial/export", params={**P, "format": "xlsx", "lang": "en"}))
    assert wb.sheetnames == ["Balance sheet", "Income statement", "Cash flow statement"]
    labels = list(_rows(wb["Income statement"]))
    assert any(label.startswith("4110") for label in labels)                 # the account tree


def test_what_the_export_refuses(books):
    api = books("ir", "IRR", IR)
    url = "/manager-reports/financial/export"
    assert api.get(url, params={"statements": "balance_sheet,horoscope"}).status_code == 422
    assert api.get(url, params={"from_date": "2026-09-01", "to_date": "2026-01-01"}).status_code == 422
    assert api.get(url, params={"format": "csv"}).status_code == 422
    assert api.get(url, params={"lang": "de"}).status_code == 422
    one = _xlsx(api.get(url, params={**P, "format": "xlsx", "statements": "cash_flow"}))
    assert one.sheetnames == ["صورت جریان‌های نقدی"]
    r = api.get(url, params={**P, "format": "xlsx"})
    assert 'filename="financial-statements-2026-01-01-2026-09-30.xlsx"' in r.headers["content-disposition"]


def test_all_currencies_is_labelled_with_the_base_currency(books):
    api = books("uk", "GBP", UK)                                      # the page's default pick is "ALL"
    url = "/manager-reports/financial/export"
    wb = _xlsx(api.get(url, params={**P, "format": "xlsx", "currency": "ALL", "statements": "balance_sheet"}))
    assert wb.active["A3"].value.endswith("All currencies, at their value in GBP")
    _pages, text = _pdf_text(api.get(url, params={**P, "format": "pdf", "currency": "ALL", "statements": "balance_sheet"}))
    assert "All currencies, at their value in GBP" in text and "ALL" not in text.replace("All", "")
    one = _xlsx(api.get(url, params={**P, "format": "xlsx", "currency": "gbp", "statements": "balance_sheet"}))
    assert one.active["A3"].value.endswith("Amounts in GBP")


def test_who_may_export():
    from app.core.auth import SessionUser
    from app.core.permissions import user_can_access
    for role, ok in (("owner", True), ("cfo", True), ("accountant", True), ("viewer", True), ("manager", False),
                     ("employee", False)):
        u = SessionUser(user_id="u", username="u", role=role, is_admin=False)
        assert user_can_access(u, "GET", "/manager-reports/financial/export") is ok, role


def test_the_export_buttons_are_wired():
    from tests.i18n_source import i18n_text
    html = open("app/static/index.html", encoding="utf-8").read()
    assert html.count('class="btn btn-secondary btn-sm mgr-statements-export"') == 2
    js = open("app/static/js/05-reports-manager.js", encoding="utf-8").read()
    assert "'/manager-reports/financial/export?'" in js and "q.set('from_date', from)" in js
    text = i18n_text()
    assert text.count("mgrStatementsPdf:") == 4 and text.count("mgrStatementsXlsx:") == 4
