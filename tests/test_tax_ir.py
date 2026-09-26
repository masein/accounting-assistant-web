"""Iranian seasonal filings (roadmap 2026-09 §3.2): TTMS report + VAT return.

Jalali seasons and their deadlines, the per-counterparty TTMS figures (مودیان
sales left out, returns netted, other currencies and incomplete parties
listed), the VAT return and its reconciliation with the report, the Excel
file, the HTTP routes and the deadline reminders.
"""
from __future__ import annotations

import io
import uuid
from datetime import date, timedelta

import jdatetime
import pytest
from sqlalchemy import select

from app.core.auth import CSRF_COOKIE, create_session_token, generate_csrf_token
from app.core.config import settings
from app.db.tenant import use_company
from app.models.company import Company
from app.models.credit_note import CreditNote
from app.models.entity import Entity
from app.models.invoice import Invoice
from app.models.invoice_item import InvoiceItem
from app.services import tax_ir
from app.services.tax_ir import Season, invoice_split, person_type, season_of

SUMMER_1405 = Season(1405, 2)          # 1405/04/01 – 1405/06/31


def _j(y, m, d):
    return jdatetime.date(y, m, d).togregorian()


# --- seasons and deadlines ------------------------------------------------------------

@pytest.mark.parametrize("season, start, end", [
    (Season(1405, 1), (1405, 1, 1), (1405, 3, 31)),
    (Season(1405, 2), (1405, 4, 1), (1405, 6, 31)),
    (Season(1405, 3), (1405, 7, 1), (1405, 9, 30)),
    (Season(1405, 4), (1405, 10, 1), (1405, 12, 29)),   # 1405 is not a leap year
    (Season(1403, 4), (1403, 10, 1), (1403, 12, 30)),   # 1403 is
])
def test_season_bounds(season, start, end):
    assert season.start == _j(*start) and season.end == _j(*end)


def test_deadlines_are_45_and_15_days_after_the_season():
    s = SUMMER_1405
    assert s.ttms_deadline == s.end + timedelta(days=45)
    assert s.vat_deadline == s.end + timedelta(days=15)
    d = s.as_dict()
    assert d["name"] == "تابستان 1405" and d["end_jalali"] == "1405/06/31"
    assert d["vat_deadline_jalali"] == "1405/07/15" and d["ttms_deadline_jalali"] == "1405/08/15"


def test_season_of_and_previous():
    assert season_of(_j(1405, 7, 4)) == Season(1405, 3)
    assert season_of(_j(1405, 6, 31)) == Season(1405, 2)
    assert Season(1405, 1).previous() == Season(1404, 4)
    assert Season(1405, 3).previous() == Season(1405, 2)
    with pytest.raises(ValueError):
        tax_ir.parse_season(2026, 2)
    with pytest.raises(ValueError):
        tax_ir.parse_season(1405, 5)


# --- building blocks ---------------------------------------------------------------------

def _inv(items, amount=None):
    inv = Invoice(number="X", kind="sales", status="issued", issue_date=date.today(), due_date=date.today(),
                  amount=amount or 0, currency="IRR")
    inv.items = [InvoiceItem(product_name=n, quantity=1, unit_price=t, line_total=t, tax_rate=r, taxable=tx,
                             tax_treatment=tr) for n, t, r, tx, tr in items]
    return inv


def test_invoice_split_by_treatment():
    s = invoice_split(_inv([("a", 1_000_000, 10, True, "standard"), ("b", 500_000, 0, False, "exempt"),
                            ("c", 200_000, 10, True, "zero_rated")]))
    assert s == {"taxable_base": 1_000_000, "exempt_base": 700_000, "base": 1_700_000, "vat": 100_000}
    assert invoice_split(_inv([], amount=42))["vat"] == 0


@pytest.mark.parametrize("nid, econ, expected", [
    ("0012345678", "", "حقیقی"), ("10101234567", "", "حقوقی"), ("", "10101234567", "حقوقی"),
    ("", "41111111111111", "حقیقی"), ("", "", "نامشخص"), ("۰۰۱۲۳۴۵۶۷۸", "", "حقیقی"),   # Persian digits count
    ("12345", "", "نامشخص"),
])
def test_person_type(nid, econ, expected):
    assert person_type(Entity(name="x", type="client", national_id=nid, economic_code=econ)) == expected


def test_identity_numbers_are_written_in_latin_digits(db, co):
    with use_company(co["company"].id):
        e = Entity(id=uuid.uuid4(), name="Persian digits", type="client", national_id="۰۰۱۲۳۴۵۶۷۸",
                   economic_code="٤١١١", postal_code="۱۲۳۴۵۶۷۸۹۰", address="x", phone="۰۲۱۸۸")
        db.add(e)
        db.commit()
    _invoice(db, co, entity=e, base=100)
    r = _report(db, co)["sales"][0]
    assert (r["national_id"], r["economic_code"], r["postal_code"], r["phone"]) == \
        ("0012345678", "4111", "1234567890", "02188")


def test_missing_fields():
    full = Entity(name="x", type="client", national_id="0012345678", economic_code="411", postal_code="1234567890",
                  address="Tehran", phone="021")
    assert tax_ir.missing_fields(full) == []
    assert tax_ir.missing_fields(Entity(name="y", type="client")) == ["national_id", "economic_code", "postal_code",
                                                                      "address", "phone"]
    assert tax_ir.missing_fields(None) == ["entity"]


# --- the report on a private company -------------------------------------------------------

@pytest.fixture()
def co(db, client):
    c = Company(id=uuid.uuid4(), name="Tax Co", slug=f"tax-{uuid.uuid4().hex[:8]}", locale="ir",
                base_currency="IRR", status="active", token_version=0)
    db.add(c)
    db.commit()
    tok = create_session_token(user_id=str(uuid.uuid4()), username="acc", is_admin=False, company_id=str(c.id),
                               role="accountant")
    client.cookies.set(settings.auth_cookie_name, tok)
    client.cookies.set(CSRF_COOKIE, generate_csrf_token())
    yield {"company": c, "client": client}
    from tests.test_admin_audit import _purge_company
    _purge_company(db, str(c.id))


def _party(db, co, name, *, complete=True, kind="client", legal=True):
    with use_company(co["company"].id):
        e = Entity(id=uuid.uuid4(), name=name, type=kind)
        if complete:
            e.national_id = "10101234567" if legal else "0012345678"
            e.economic_code = "10101234567"
            e.postal_code = "1234567890"
            e.address = "Tehran"
            e.phone = "02188888888"
            e.province, e.city = "Tehran", "Tehran"
        db.add(e)
        db.commit()
    return e


def _invoice(db, co, *, kind="sales", entity=None, day=(1405, 5, 10), base=1_000_000, rate=10, status="issued",
             currency="IRR", moadian=None, exempt=0):
    with use_company(co["company"].id):
        d = _j(*day)
        items = [InvoiceItem(product_name="svc", quantity=1, unit_price=base, line_total=base, tax_rate=rate,
                             taxable=rate > 0)]
        if exempt:
            items.append(InvoiceItem(product_name="bread", quantity=1, unit_price=exempt, line_total=exempt,
                                     tax_rate=0, taxable=False, tax_treatment="exempt"))
        vat = base * rate // 100
        inv = Invoice(id=uuid.uuid4(), number=f"INV-{uuid.uuid4().hex[:6]}", kind=kind, status=status,
                      issue_date=d, due_date=d, amount=base + exempt + vat, currency=currency,
                      entity_id=entity.id if entity else None, moadian_status=moadian)
        inv.items = items
        db.add(inv)
        db.commit()
    return inv


def _credit(db, co, inv, amount, day=(1405, 6, 1)):
    with use_company(co["company"].id):
        cn = CreditNote(invoice_id=inv.id, entity_id=inv.entity_id, kind=inv.kind, date=_j(*day), amount=amount,
                        currency="IRR", note_type="reduction")
        db.add(cn)
        db.commit()
    return cn


def _report(db, co, **kw):
    with use_company(co["company"].id):
        return tax_ir.quarterly_report(db, SUMMER_1405, **kw)


def test_quarterly_report_groups_per_counterparty(db, co):
    a = _party(db, co, "Alpha Ltd")
    b = _party(db, co, "Beta", complete=False)
    s = _party(db, co, "Supplier S", kind="supplier", legal=False)
    _invoice(db, co, entity=a, base=1_000_000)
    _invoice(db, co, entity=a, base=2_000_000, exempt=500_000)
    _invoice(db, co, entity=b, base=300_000)
    _invoice(db, co, kind="purchase", entity=s, base=400_000)
    _invoice(db, co, entity=a, base=9_999_999, status="draft")                     # not recognised
    _invoice(db, co, entity=a, base=8_888_888, day=(1405, 7, 1))                  # next season
    rep = _report(db, co)
    sales = {r["name"]: r for r in rep["sales"]}
    assert sales["Alpha Ltd"]["invoice_count"] == 2
    assert sales["Alpha Ltd"]["base"] == 3_500_000 and sales["Alpha Ltd"]["vat"] == 300_000
    assert sales["Alpha Ltd"]["total"] == 3_800_000 and sales["Alpha Ltd"]["person_type"] == "حقوقی"
    assert sales["Beta"]["missing"] == ["national_id", "economic_code", "postal_code", "address", "phone"]
    [purchase] = rep["purchases"]
    assert purchase["name"] == "Supplier S" and purchase["vat"] == 40_000 and purchase["person_type"] == "حقیقی"
    assert rep["totals"]["sales"]["base"] == 3_800_000 and rep["totals"]["sales"]["vat"] == 330_000
    assert {i["name"] for i in rep["incomplete"]} == {"Beta"}
    assert rep["ready"] is False
    assert all(d["issue_date_jalali"].startswith("1405/05") for d in rep["details"])


def test_moadian_sales_are_left_out_unless_included(db, co):
    a = _party(db, co, "Alpha Ltd")
    _invoice(db, co, entity=a, base=1_000_000, moadian="confirmed")
    _invoice(db, co, entity=a, base=500_000, moadian="exported")                  # sent, not accepted: still reported
    rep = _report(db, co)
    assert rep["totals"]["sales"]["base"] == 500_000
    assert rep["excluded_moadian"] == {"count": 1, "base": 1_000_000, "vat": 100_000}
    assert _report(db, co, include_moadian=True)["totals"]["sales"]["base"] == 1_500_000


def test_returns_are_netted_with_their_share_of_vat(db, co):
    a = _party(db, co, "Alpha Ltd")
    inv = _invoice(db, co, entity=a, base=1_000_000)            # gross 1,100,000
    _credit(db, co, inv, 110_000)                               # a tenth back: 100,000 + 10,000 VAT
    r = _report(db, co)["sales"][0]
    assert (r["returns_base"], r["returns_vat"]) == (100_000, 10_000)
    assert (r["net_base"], r["net_vat"]) == (900_000, 90_000)


def test_other_currencies_are_listed_not_summed(db, co):
    a = _party(db, co, "Alpha Ltd")
    _invoice(db, co, entity=a, base=5_000, currency="USD")
    rep = _report(db, co)
    assert rep["sales"] == [] and len(rep["other_currency"]) == 1 and rep["ready"] is False


def test_vat_return_and_reconciliation(db, co):
    a = _party(db, co, "Alpha Ltd")
    s = _party(db, co, "Supplier S", kind="supplier")
    inv = _invoice(db, co, entity=a, base=2_000_000, exempt=300_000)
    _invoice(db, co, entity=a, base=1_000_000, moadian="confirmed")              # still in the return
    _invoice(db, co, kind="purchase", entity=s, base=500_000)
    _credit(db, co, inv, 220_000)
    with use_company(co["company"].id):
        ret = tax_ir.vat_return(db, SUMMER_1405)
        rec = tax_ir.reconciliation(db, SUMMER_1405)
    assert ret["sales"]["taxable_base"] == 3_000_000 and ret["sales"]["exempt_base"] == 300_000
    assert ret["sales"]["vat"] == 300_000
    # the credit is 220,000 of a 2,500,000 gross invoice carrying 200,000 VAT → 17,600 VAT back
    assert ret["sales"]["returns_vat"] == 17_600 and ret["sales"]["net_vat"] == 282_400
    assert ret["purchases"]["net_vat"] == 50_000
    assert ret["payable"] == 232_400 and ret["credit_carried_forward"] == 0 and ret["net"] == 232_400
    assert rec["matches"] is True and all(r["matches"] for r in rec["rows"])


def test_more_input_than_output_is_a_credit(db, co):
    s = _party(db, co, "Supplier S", kind="supplier")
    _invoice(db, co, kind="purchase", entity=s, base=900_000)
    with use_company(co["company"].id):
        ret = tax_ir.vat_return(db, SUMMER_1405)
    assert ret["payable"] == 0 and ret["credit_carried_forward"] == 90_000


def test_another_companys_invoices_never_appear(db, co):
    _invoice(db, co, entity=_party(db, co, "Mine"), base=1_000)
    other = Company(id=uuid.uuid4(), name="Other", slug=f"o-{uuid.uuid4().hex[:8]}", locale="ir",
                    base_currency="IRR", status="active", token_version=0)
    db.add(other)
    db.commit()
    try:
        _invoice(db, {"company": other}, entity=_party(db, {"company": other}, "Theirs"), base=7_777_777)
        names = {r["name"] for r in _report(db, co)["sales"]}
        assert names == {"Mine"}
    finally:
        from tests.test_admin_audit import _purge_company
        _purge_company(db, str(other.id))


# --- HTTP and the Excel file ------------------------------------------------------------------

def test_routes_and_the_excel_workbook(db, co):
    from openpyxl import load_workbook
    a = _party(db, co, "Alpha Ltd")
    b = _party(db, co, "Beta", complete=False)
    _invoice(db, co, entity=a, base=1_000_000)
    _invoice(db, co, entity=b, base=200_000)
    c = co["client"]
    assert c.get("/tax/ir/seasons").status_code == 200
    rep = c.get("/tax/ir/quarterly?year=1405&season=2").json()
    assert rep["season"]["name"] == "تابستان 1405" and len(rep["sales"]) == 2
    ret = c.get("/tax/ir/vat-return?year=1405&season=2").json()
    assert ret["output_vat"] == 120_000 and ret["reconciliation"]["matches"] is True
    r = c.get("/tax/ir/quarterly/export?year=1405&season=2")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert 'filename="ttms-1405-2.xlsx"' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["خلاصه", "فروش", "خرید", "صورتحساب‌ها", "اطلاعات ناقص"]
    sales = wb["فروش"]
    assert sales.sheet_view.rightToLeft is True
    headers = [cell.value for cell in sales[1]]
    assert headers[:5] == ["ردیف", "نوع شخص", "شناسه/کد ملی", "کد اقتصادی", "نام"]
    rows = {row[4]: row for row in sales.iter_rows(min_row=2, values_only=True)}
    assert rows["Alpha Ltd"][11] == 1_000_000 and rows["Alpha Ltd"][12] == 100_000
    assert "کد پستی" in rows["Beta"][-1]
    assert [r[1] for r in wb["اطلاعات ناقص"].iter_rows(min_row=2, values_only=True)] == ["Beta"]
    summary = {r[0]: r[1] for r in wb["خلاصه"].iter_rows(values_only=True) if r[0]}
    assert summary["اظهارنامه — مالیات فروش (ریال)"] == 120_000
    assert summary["تطبیق گزارش فصلی با اظهارنامه"] == "مطابق"


def test_bad_season_and_who_may_read(db, co, client):
    assert co["client"].get("/tax/ir/quarterly?year=1405&season=9").status_code == 422
    assert co["client"].get("/tax/ir/quarterly?year=2026&season=1").status_code == 422
    tok = create_session_token(user_id=str(uuid.uuid4()), username="emp", is_admin=False,
                               company_id=str(co["company"].id), role="employee")
    client.cookies.set(settings.auth_cookie_name, tok)
    assert client.get("/tax/ir/quarterly?year=1405&season=2").status_code == 403


# --- reminders -----------------------------------------------------------------------------------

def _tax_notes(db, co):
    from app.models.notification import Notification
    with use_company(co["company"].id):
        return {n.dedupe_key: n for n in db.execute(select(Notification).where(
            Notification.kind == "tax_filing", Notification.dismissed_at.is_(None))).scalars()}


def _refresh(db, co, today):
    from app.services.notification_service import refresh_notifications
    with use_company(co["company"].id):
        refresh_notifications(db, today=today)
        db.commit()


def test_deadline_reminders_for_the_season_just_ended(db, co):
    _invoice(db, co, entity=_party(db, co, "Alpha Ltd"))
    s = SUMMER_1405
    _refresh(db, co, s.vat_deadline - timedelta(days=5))
    notes = _tax_notes(db, co)
    assert set(notes) == {"tax-vat-1405-2"} and notes["tax-vat-1405-2"].level == "warning"
    _refresh(db, co, s.vat_deadline - timedelta(days=1))
    assert _tax_notes(db, co)["tax-vat-1405-2"].level == "high"
    _refresh(db, co, s.ttms_deadline - timedelta(days=8))
    notes = _tax_notes(db, co)
    assert set(notes) == {"tax-ttms-1405-2"} and "TTMS" in notes["tax-ttms-1405-2"].title
    _refresh(db, co, s.ttms_deadline + timedelta(days=1))
    assert _tax_notes(db, co) == {}


def test_no_reminders_without_trade_or_outside_iran(db, co):
    s = SUMMER_1405
    _refresh(db, co, s.vat_deadline - timedelta(days=3))
    assert _tax_notes(db, co) == {}                      # no invoices in the season
    _invoice(db, co, entity=_party(db, co, "Alpha Ltd"))
    co["company"].locale = "uk"
    db.commit()
    _refresh(db, co, s.vat_deadline - timedelta(days=3))
    assert _tax_notes(db, co) == {}
