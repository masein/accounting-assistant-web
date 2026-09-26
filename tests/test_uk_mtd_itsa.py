"""UK Making Tax Digital for Income Tax (roadmap 2026-09 §3.6, part two).

Quarterly figures in HMRC's categories from a UK company's ledger —
self-employment and UK property, cumulative from April, disallowable items,
overrides, consolidated expenses, the calendar basis, the mandation test, the
HMRC request bodies, the workbook and the reminder."""
from __future__ import annotations

import io
import json
import uuid
from datetime import date, datetime, timedelta, timezone

import pytest
from sqlalchemy import select

from app.core.auth import CSRF_COOKIE, create_session_token, generate_csrf_token
from app.core.config import settings
from app.db.seed import seed_chart_if_empty
from app.db.tenant import use_company
from app.models.account import Account
from app.models.company import Company
from app.models.transaction import Transaction, TransactionLine
from app.services.uk_mtd import categories as C
from app.services.uk_mtd import itsa
from app.services.uk_mtd import periods as P
from app.services.uk_mtd import settings as S


@pytest.fixture()
def books(db, client):
    c = Company(id=uuid.uuid4(), name="Sole Trader", slug=f"st-{uuid.uuid4().hex[:8]}", locale="uk",
                base_currency="GBP", status="active", token_version=0)
    db.add(c)
    db.commit()
    with use_company(c.id):
        seed_chart_if_empty(db, locale="uk")
        db.commit()
        acc = {a.code: a for a in db.execute(select(Account)).scalars()}

    def login(role="owner"):
        tok = create_session_token(user_id=str(uuid.uuid4()), username=role, is_admin=(role == "owner"),
                                   company_id=str(c.id), role=role)
        csrf = generate_csrf_token()
        client.cookies.set(settings.auth_cookie_name, tok)
        client.cookies.set(CSRF_COOKIE, csrf)
        from tests.conftest import _CSRFTestClient
        return _CSRFTestClient(client, csrf)
    yield {"company": c, "acc": acc, "login": login}
    from tests.test_admin_audit import _purge_company
    _purge_company(db, str(c.id))


BANK = "1200"


def _post(db, b, day, debit, credit, amount, *, currency="GBP", deleted=False):
    with use_company(b["company"].id):
        t = Transaction(id=uuid.uuid4(), date=day, reference="J", description="j", currency=currency)
        if deleted:
            t.deleted_at = datetime.now(timezone.utc)
        db.add(t)
        db.flush()
        db.add_all([TransactionLine(transaction_id=t.id, account_id=b["acc"][debit].id, debit=amount, credit=0),
                    TransactionLine(transaction_id=t.id, account_id=b["acc"][credit].id, debit=0, credit=amount)])
        db.commit()


def _seed_year(db, b):
    q1 = date(2026, 5, 10)
    _post(db, b, q1, BANK, "4000", 10_000)           # sales
    _post(db, b, q1, "5000", BANK, 3_000)            # purchases
    _post(db, b, q1, "7200", BANK, 1_200)            # rent
    _post(db, b, q1, "7500", BANK, 500)              # travel and entertainment
    _post(db, b, q1, "8500", BANK, 800)              # depreciation
    _post(db, b, q1, "8000", BANK, 30)               # bank charges
    _post(db, b, q1, BANK, "8300", 50)               # interest received: not trading income
    _post(db, b, q1, "9000", BANK, 999)              # tax: never in the update
    _post(db, b, q1, "5000", BANK, 7_777, deleted=True)
    _post(db, b, q1, BANK, "4000", 4_444, currency="USD")
    q2 = date(2026, 8, 20)
    _post(db, b, q2, BANK, "4000", 5_000)
    _post(db, b, q2, "7100", BANK, 2_000)            # wages


def _update(db, b, quarter, *, source="self_employment", overrides=None, basis="standard"):
    with use_company(b["company"].id):
        return itsa.quarterly_update(db, P.ItsaQuarter(2026, quarter, basis), source=source,
                                     overrides=overrides or {}, currency="GBP")


def _cats(lines):
    return {l["category"]: (l["quarter"], l["year_to_date"]) for l in lines if l["quarter"] or l["year_to_date"]}


# --- catalogue and defaults ------------------------------------------------------------------

def test_catalogue_uses_hmrc_field_names():
    assert len(C.SELF_EMPLOYMENT_EXPENSES) == 15 and "consolidatedExpenses" not in C.SELF_EMPLOYMENT_EXPENSES
    assert C.kind_of("turnover", "self_employment") == "income"
    assert C.kind_of("residentialFinancialCost", "uk_property") == "expenses"
    assert C.kind_of("turnover", "uk_property") is None


@pytest.mark.parametrize("code, source, expected", [
    ("4000", "self_employment", "turnover"), ("7500", "self_employment", "carVanTravelExpenses"),
    ("8100", "self_employment", "interestOnBankOtherLoans"), ("8300", "self_employment", C.EXCLUDED),
    ("9100", "self_employment", C.EXCLUDED), ("6123", "self_employment", "advertisingCosts"),
    ("4000", "uk_property", "periodAmount"), ("8100", "uk_property", "residentialFinancialCost"),
    ("8500", "uk_property", C.EXCLUDED), ("7700", "uk_property", "repairsAndMaintenance"),
])
def test_default_mapping(code, source, expected):
    assert C.default_category(code, source) == expected


# --- self-employment ------------------------------------------------------------------------------

def test_self_employment_quarter_one(db, books):
    _seed_year(db, books)
    up = _update(db, books, 1)
    assert _cats(up["income"]) == {"turnover": (10_000, 10_000)}
    assert _cats(up["expenses"]) == {"costOfGoods": (3_000, 3_000), "premisesRunningCosts": (1_200, 1_200),
                                     "carVanTravelExpenses": (500, 500), "depreciation": (800, 800),
                                     "financeCharges": (30, 30)}
    assert up["totals"]["quarter"]["profit"] == 10_000 - 5_530
    assert {e["code"] for e in up["excluded"]} == {"8300", "9000"}
    body = up["hmrc_body"]
    assert body["periodDates"] == {"periodStartDate": "2026-04-06", "periodEndDate": "2026-07-05"}
    assert body["periodIncome"] == {"turnover": 10000.0}
    assert body["periodExpenses"]["costOfGoods"] == 3000.0 and "consolidatedExpenses" not in body["periodExpenses"]
    assert body["periodDisallowableExpenses"] == {"depreciationDisallowable": 800.0}


def test_updates_are_cumulative_from_april(db, books):
    _seed_year(db, books)
    up = _update(db, books, 2)
    inc, exp = _cats(up["income"]), _cats(up["expenses"])
    assert inc["turnover"] == (5_000, 15_000)
    assert exp["wagesAndStaffCosts"] == (2_000, 2_000) and exp["costOfGoods"] == (0, 3_000)
    assert up["hmrc_body"]["periodDates"] == {"periodStartDate": "2026-04-06", "periodEndDate": "2026-10-05"}
    assert up["hmrc_body"]["periodIncome"]["turnover"] == 15000.0


def test_consolidated_expenses_below_the_vat_threshold(db, books):
    _seed_year(db, books)
    up = _update(db, books, 2)
    assert up["consolidated_allowed"] is True
    assert up["hmrc_body_consolidated"]["periodExpenses"] == {"consolidatedExpenses": float(5_530 + 2_000)}
    _post(db, books, date(2026, 9, 1), BANK, "4000", 100_000)
    up = _update(db, books, 2)
    assert up["consolidated_allowed"] is False and up["hmrc_body_consolidated"] is None


def test_overrides_move_an_account(db, books):
    _seed_year(db, books)
    up = _update(db, books, 1, overrides={"7500": "businessEntertainmentCosts"})
    exp = _cats(up["expenses"])
    assert "carVanTravelExpenses" not in exp and exp["businessEntertainmentCosts"] == (500, 500)
    assert up["hmrc_body"]["periodDisallowableExpenses"] == {"businessEntertainmentCostsDisallowable": 500.0,
                                                            "depreciationDisallowable": 800.0}
    row = next(a for a in up["accounts"] if a["code"] == "7500")
    assert row["overridden"] is True
    # an account mapped to a category of the other source is reported as unmapped
    up = _update(db, books, 1, overrides={"7200": "residentialFinancialCost"})
    assert [u["code"] for u in up["unmapped"]] == ["7200"]


def test_calendar_basis(db, books):
    _post(db, books, date(2026, 4, 3), BANK, "4000", 700)   # before 6 April: calendar Q1 only
    assert _cats(_update(db, books, 1, basis="calendar")["income"]) == {"turnover": (700, 700)}
    assert _update(db, books, 1)["income"][0]["year_to_date"] == 0


# --- UK property ------------------------------------------------------------------------------------

def test_uk_property_update(db, books):
    q1 = date(2026, 5, 1)
    _post(db, books, q1, BANK, "4000", 6_000)        # rents
    _post(db, books, q1, "8100", BANK, 900)          # mortgage interest
    _post(db, books, q1, "7700", BANK, 400)          # repairs
    _post(db, books, q1, "8500", BANK, 300)          # depreciation: never allowable
    up = _update(db, books, 1, source="uk_property")
    assert _cats(up["income"]) == {"periodAmount": (6_000, 6_000)}
    assert _cats(up["expenses"]) == {"residentialFinancialCost": (900, 900), "repairsAndMaintenance": (400, 400)}
    assert {e["code"] for e in up["excluded"]} == {"8500"}
    body = up["hmrc_body"]
    assert body["fromDate"] == "2026-04-06" and body["toDate"] == "2026-07-05"
    assert body["ukProperty"]["income"] == {"periodAmount": 6000.0}
    assert body["ukProperty"]["expenses"] == {"repairsAndMaintenance": 400.0, "residentialFinancialCost": 900.0}
    cons = up["hmrc_body_consolidated"]["ukProperty"]["expenses"]
    assert cons == {"consolidatedExpenses": 400.0, "residentialFinancialCost": 900.0}


def test_rent_a_room_nests_like_hmrc(db, books):
    _post(db, books, date(2026, 5, 1), BANK, "4200", 1_500)
    up = _update(db, books, 1, source="uk_property", overrides={"4200": "rentARoomRents"})
    assert up["hmrc_body"]["ukProperty"]["income"] == {"rentARoom": {"rentsReceived": 1500.0}}


def test_unknown_source_is_refused(db, books):
    with pytest.raises(ValueError):
        _update(db, books, 1, source="none")


# --- mandation ---------------------------------------------------------------------------------------

def test_mandation_uses_income_two_years_earlier(db, books):
    _post(db, books, date(2024, 10, 1), BANK, "4000", 60_000)   # 2024-25
    _post(db, books, date(2025, 10, 1), BANK, "4000", 25_000)   # 2025-26
    with use_company(books["company"].id):
        m26 = itsa.mandation(db, 2026, source="self_employment")
        m27 = itsa.mandation(db, 2027, source="self_employment")
    assert (m26["based_on"], m26["qualifying_income"], m26["threshold"], m26["required"]) == ("2024-25", 60_000, 50_000, True)
    assert (m27["based_on"], m27["qualifying_income"], m27["threshold"], m27["required"]) == ("2025-26", 25_000, 30_000, False)


# --- HTTP ------------------------------------------------------------------------------------------------

def test_routes_workbook_and_settings(db, books):
    from openpyxl import load_workbook
    _seed_year(db, books)
    owner = books["login"]("owner")
    assert owner.get("/tax/uk/itsa/update?tax_year=2026-27&quarter=1").status_code == 409   # no source yet
    assert owner.put("/tax/uk/settings", json={"income_source": "self_employment"}).status_code == 200
    q = owner.get("/tax/uk/itsa/quarters?tax_year=2026-27").json()
    assert [x["deadline"] for x in q["quarters"]] == ["2026-08-07", "2026-11-07", "2027-02-07", "2027-05-07"]
    assert q["mandation"]["tax_year"] == "2026-27"
    up = owner.get("/tax/uk/itsa/update?tax_year=2026-27&quarter=1").json()
    assert up["hmrc_body"]["periodIncome"] == {"turnover": 10000.0}
    assert owner.get("/tax/uk/itsa/update?tax_year=2026-28&quarter=1").status_code == 422
    cats = owner.get("/tax/uk/itsa/categories").json()
    assert cats["source"] == "self_employment"
    row = next(a for a in cats["accounts"] if a["code"] == "7500")
    assert row["default"] == "carVanTravelExpenses" and row["override"] is None
    assert owner.put("/tax/uk/settings", json={"category_overrides": {"7500": "nonsense"}}).status_code == 422
    assert owner.put("/tax/uk/settings", json={"category_overrides": {"7500": "businessEntertainmentCosts"}}).status_code == 200
    up = owner.get("/tax/uk/itsa/update?tax_year=2026-27&quarter=1").json()
    assert "businessEntertainmentCostsDisallowable" in up["hmrc_body"]["periodDisallowableExpenses"]
    r = owner.get("/tax/uk/itsa/update/export?tax_year=2026-27&quarter=1")
    assert r.status_code == 200 and 'filename="mtd-itsa-2026-27-q1.xlsx"' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Update", "Accounts", "HMRC body"]
    body = json.loads(wb["HMRC body"]["A2"].value)
    assert body["periodIncome"]["turnover"] == 10000.0
    codes = {row[0] for row in wb["Accounts"].iter_rows(min_row=2, values_only=True)}
    assert {"4000", "7500", "8300"} <= codes
    assert books["login"]("employee").get("/tax/uk/itsa/update?tax_year=2026-27&quarter=1").status_code == 403


# --- reminder --------------------------------------------------------------------------------------------

def test_quarterly_update_reminder(db, books):
    from app.models.notification import Notification
    from app.services.notification_service import refresh_notifications

    def notes(today):
        with use_company(books["company"].id):
            refresh_notifications(db, today=today)
            db.commit()
            return {n.dedupe_key: n for n in db.execute(select(Notification).where(
                Notification.kind == "tax_filing", Notification.dismissed_at.is_(None))).scalars()}
    due = P.ItsaQuarter(2026, 1, "standard").deadline                     # 7 Aug 2026
    assert notes(due - timedelta(days=5)) == {}                          # not keeping MTD records
    with use_company(books["company"].id):
        S.save_settings(db, income_source="uk_property")
        db.commit()
    got = notes(due - timedelta(days=5))
    assert set(got) == {"uk-itsa-2026-1"} and got["uk-itsa-2026-1"].level == "warning"
    assert notes(due)["uk-itsa-2026-1"].level == "high"
    assert notes(due + timedelta(days=1)) == {}
