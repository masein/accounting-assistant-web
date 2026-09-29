"""A viewer reads the reports but not bank or identity numbers — in documents
too. The JSON entity reads already stripped them; a viewer's statement PDF
still printed the party's national id, and the close pack carried the bank
statement lines and the bank's account number."""
from __future__ import annotations

import io
from datetime import date

import pytest

from tests.test_close_pack import URL as PACK, UK, _statement
from tests.test_statement_export import _company


def _as(client, role, cid):
    from app.core.auth import CSRF_COOKIE, create_session_token, generate_csrf_token
    from app.core.config import settings
    from tests.conftest import _CSRFTestClient
    import uuid
    csrf = generate_csrf_token()
    client.cookies.clear()
    client.cookies.set(settings.auth_cookie_name, create_session_token(
        user_id=str(uuid.uuid4()), username=role, is_admin=role == "owner", role=role, company_id=cid))
    client.cookies.set(CSRF_COOKIE, csrf)
    return _CSRFTestClient(client, csrf)


@pytest.fixture()
def books(client, db):
    api, cid = _company(client, db, "uk", "GBP")
    for d, desc, dr, cr, amount in UK:
        r = api.post("/transactions", json={"date": d, "description": desc, "currency": "GBP", "lines": [
            {"account_code": dr, "debit": amount, "credit": 0}, {"account_code": cr, "debit": 0, "credit": amount}]})
        assert r.status_code == 201, r.text
    yield client, cid
    from tests.test_admin_audit import _purge_company
    client.cookies.clear()
    _purge_company(db, cid)


def _text(pdf: bytes) -> str:
    from pypdf import PdfReader
    return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(pdf)).pages)


def test_a_viewers_statement_pdf_leaves_out_the_identity_numbers(books):
    client, cid = books
    owner = _as(client, "owner", cid)
    sup = owner.post("/entities", json={"type": "supplier", "name": "Paper Ltd", "national_id": "AB123456C",
                                        "iban": "GB33BUKB20201555555555"}).json()
    url = f"/entities/{sup['id']}/statement.pdf"
    r = owner.get(url)
    assert r.status_code == 200 and "AB123456C" in _text(r.content)
    viewer = _as(client, "viewer", cid)
    assert viewer.get(f"/entities/{sup['id']}").json()["national_id"] is None       # the JSON already hid it
    r = viewer.get(url)
    assert r.status_code == 200 and "Paper Ltd" in _text(r.content)
    assert "AB123456C" not in _text(r.content)                                        # …and now the PDF does
    # the saved record is untouched
    assert _as(client, "owner", cid).get(f"/entities/{sup['id']}").json()["national_id"] == "AB123456C"


def test_a_viewers_close_pack_leaves_out_the_bank(books, db):
    from openpyxl import load_workbook
    client, cid = books
    _statement(db, cid, [(date(2026, 8, 10), "Card payment", 0, 40_000, 130_000, "matched"),
                         (date(2026, 8, 20), "Unknown transfer", 250, 0, 129_750, "unmatched")])
    sheets = {}
    for role in ("owner", "accountant", "viewer"):
        r = _as(client, role, cid).get(PACK, params={"month": "2026-08", "format": "xlsx"})
        assert r.status_code == 200, (role, r.text)
        sheets[role] = load_workbook(io.BytesIO(r.content)).sheetnames
    assert "Bank reconciliation" in sheets["owner"] and "Bank lines not matched" in sheets["accountant"]
    assert "Bank reconciliation" not in sheets["viewer"] and "Bank lines not matched" not in sheets["viewer"]
    assert "Trial balance" in sheets["viewer"] and "Budget vs actual" in sheets["viewer"]      # the rest is there
    r = _as(client, "viewer", cid).get(PACK, params={"month": "2026-08", "format": "pdf"})
    text = _text(r.content)
    assert "12345678" not in text and "Unknown transfer" not in text
    # the checklist still says how far the matching got — a count, not the lines
    chk = _as(client, "viewer", cid).get(f"{PACK}/checklist", params={"month": "2026-08"}).json()
    assert {i["key"]: i["state"] for i in chk["items"]}["bank"] == "warn"


def test_invoice_and_quote_pdfs_and_the_ttms_figures_too(client, db):
    from tests.test_statement_export import _company
    from tests.test_admin_audit import _purge_company
    import uuid as _uuid
    owner, cid = _company(client, db, "ir", "IRR")
    try:
        owner = _as(client, "owner", cid)
        buyer = owner.post("/entities", json={"type": "client", "name": "Aria Trading", "national_id": "0012345678",
                                              "economic_code": "411111111111"}).json()
        inv = owner.post("/invoices", json={"number": f"V-{_uuid.uuid4().hex[:4]}", "kind": "sales",
                                            "status": "issued", "issue_date": "2026-07-01", "due_date": "2026-07-30",
                                            "amount": 5_000_000, "currency": "IRR", "entity_id": buyer["id"]}).json()
        q = owner.post("/quotes", json={"entity_id": buyer["id"], "issue_date": "2026-07-01",
                                        "valid_until": "2026-07-31", "amount": 1_000_000})
        assert q.status_code == 201, q.text

        def shows_id(pdf: bytes) -> bool:
            # the national-id line of the party card ("شناسه ملی: …"; its Persian digits don't
            # extract from the PDF, the label does — and the line is left out when the id is)
            return "شناسه ملی" in _text(pdf)

        for url in (f"/invoices/{inv['id']}/pdf", f"/quotes/{q.json()['id']}/pdf"):
            assert shows_id(_as(client, "owner", cid).get(url).content), url
            r = _as(client, "viewer", cid).get(url)
            assert r.status_code == 200 and not shows_id(r.content), url
        # TTMS: the viewer sees the figures and what's missing, not the national id
        params = {"year": 1405, "season": 2}
        mine = _as(client, "owner", cid).get("/tax/ir/quarterly", params=params).json()
        theirs = _as(client, "viewer", cid).get("/tax/ir/quarterly", params=params).json()

        def rows(report):
            return [r for key in ("sales", "purchases") for r in report.get(key, []) if isinstance(r, dict)]

        owner_rows, viewer_rows = rows(mine), rows(theirs)
        assert any(r.get("national_id") == "0012345678" for r in owner_rows)
        assert viewer_rows and all(not r.get("national_id") for r in viewer_rows)
        assert [r.get("total") for r in viewer_rows] == [r.get("total") for r in owner_rows]
        assert [r.get("economic_code") for r in viewer_rows] == [r.get("economic_code") for r in owner_rows]
    finally:
        client.cookies.clear()
        _purge_company(db, cid)


def test_the_rule_lives_in_one_place():
    from app.core.redaction import IDENTITY_FIELDS
    assert set(IDENTITY_FIELDS) == {"account_number", "iban", "sort_code", "account_holder", "national_id"}
    import inspect

    from app.api import entities, invoices, quotes
    from app.services import tax_ir
    from app.services.ai_accountant import read_tools
    from app.services.reporting import close_pack
    for mod in (entities, invoices, quotes, tax_ir, close_pack, read_tools):
        src = inspect.getsource(mod)
        assert "app.core.redaction" in src, mod.__name__


def _entity(**kw):
    from types import SimpleNamespace
    base = dict(phone="0912", email=None, address="Tehran", tax_id=None, economic_code="411111111111",
                national_id="0012345678", contact_person=None, bank_name="Mellat", account_holder="Aria",
                account_number="12345678", iban="IR820540102680020817909002", sort_code="20-20-15")
    return SimpleNamespace(**{**base, **kw})


@pytest.mark.parametrize("role, shown", [("owner", True), ("accountant", True), ("personal", True),
                                         ("viewer", False), ("manager", False), ("employee", False)])
def test_the_chat_tools_follow_the_same_rule(role, shown):
    """find_entity / list_entities hand the model a party's details — the bank and
    identity numbers only when the caller could read them on the page."""
    from types import SimpleNamespace

    from app.core.redaction import IDENTITY_FIELDS
    from app.core.request_context import clear_current_user, set_current_user
    from app.services.ai_accountant.read_tools import _entity_details
    set_current_user(SimpleNamespace(user_id="u", username=role, role=role, entity_id=None, is_superadmin=False))
    try:
        out = _entity_details(_entity())
    finally:
        clear_current_user()
    assert out["phone"] == "0912" and out["economic_code"] == "411111111111" and out["bank_name"] == "Mellat"
    assert all((f in out) is shown for f in IDENTITY_FIELDS), out
    assert _entity_details(_entity())["iban"].startswith("IR82")         # a job (no caller) keeps them


def test_every_role_that_can_chat_can_read_bank_details():
    """Today only roles with bank:read reach the assistant (web chat and the
    messenger bot check the same route). If that ever changes, the chat tools
    above already hide the numbers — this pins the assumption so it's a choice."""
    from types import SimpleNamespace

    from app.core.permissions import ROLE_PERMISSIONS, Perm, role_can, user_can_access
    chat = sorted(r for r in ROLE_PERMISSIONS
                  if user_can_access(SimpleNamespace(role=r, is_superadmin=False), "POST", "/ai-accountant/chat"))
    assert chat and all(role_can(r, Perm.BANK_READ) for r in chat), chat
